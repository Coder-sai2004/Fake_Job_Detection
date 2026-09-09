import docx
import os

friend_path = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\INTERNSHIP REPORT.docx'
user_path = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report.docx'

doc_f = docx.Document(friend_path)
doc_u = docx.Document(user_path)

print("=== FRIEND DOCUMENT DETAILS ===")
print("Sections count:", len(doc_f.sections))
for i, s in enumerate(doc_f.sections):
    print(f"Section {i}: W={s.page_width.inches} H={s.page_height.inches} T={s.top_margin.inches} B={s.bottom_margin.inches} L={s.left_margin.inches} R={s.right_margin.inches}")
    print(f"  Header ({len(s.header.paragraphs)} p):", [p.text for p in s.header.paragraphs if p.text.strip()])
    print(f"  Footer ({len(s.footer.paragraphs)} p):", [p.text for p in s.footer.paragraphs if p.text.strip()])

print("\n=== USER DOCUMENT DETAILS ===")
print("Sections count:", len(doc_u.sections))
for i, s in enumerate(doc_u.sections):
    print(f"Section {i}: W={s.page_width.inches} H={s.page_height.inches} T={s.top_margin.inches} B={s.bottom_margin.inches} L={s.left_margin.inches} R={s.right_margin.inches}")
    print(f"  Header ({len(s.header.paragraphs)} p):", [p.text for p in s.header.paragraphs if p.text.strip()])
    print(f"  Footer ({len(s.footer.paragraphs)} p):", [p.text for p in s.footer.paragraphs if p.text.strip()])

print("\n=== FRIEND STYLES & SIZES USED ===")
styles_used = {}
for p in doc_f.paragraphs:
    if p.text.strip():
        st = p.style.name
        sizes = [r.font.size.pt for r in p.runs if r.font.size]
        fonts = [r.font.name for r in p.runs if r.font.name]
        bolds = [r.bold for r in p.runs]
        key = (st, p.alignment, tuple(set(sizes)), tuple(set(fonts)), tuple(set(bolds)))
        styles_used[key] = styles_used.get(key, 0) + 1

for k, v in sorted(styles_used.items(), key=lambda x: x[1], reverse=True):
    print(f"Count {v:3d}: Style={k[0]}, Align={k[1]}, Sizes={k[2]}, Fonts={k[3]}, Bold={k[4]}")
