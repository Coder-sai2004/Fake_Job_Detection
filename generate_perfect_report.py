import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

MEDIA_DIR = "report_assets"
OUTPUT_FILE = r"C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report_Formatted.docx"
FINAL_REPLACE_FILE = r"C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report.docx"

doc = docx.Document()

# ----------------- SECTION 1: PRELIMINARY PAGES -----------------
sec1 = doc.sections[0]
sec1.page_width = Inches(8.268)   # A4 (210 mm)
sec1.page_height = Inches(11.693) # A4 (297 mm)
sec1.top_margin = Inches(1.0)
sec1.bottom_margin = Inches(1.0)
sec1.left_margin = Inches(1.0)
sec1.right_margin = Inches(1.0)
sec1.header_distance = Inches(0.49)
sec1.footer_distance = Inches(0.49)
sec1.different_first_page_header_footer = False

# Ensure Section 1 footer has NO page number
sec1.footer.is_linked_to_previous = False
for p in sec1.footer.paragraphs:
    p.text = ""

FONT_NAME = 'Times New Roman'
COLOR_BLACK = RGBColor(0, 0, 0)

def set_run_font(run, name=FONT_NAME, size_pt=12, bold=False, italic=False, color=COLOR_BLACK):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, italic=False, 
          space_before=0, space_after=6, line_spacing=1.15, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.keep_with_next = keep_with_next
    if text:
        r = p.add_run(text)
        set_run_font(r, size_pt=size_pt, bold=bold, italic=italic)
    return p

def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run()
    run.add_break(docx.enum.text.WD_BREAK.PAGE)

def add_chapter_heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size_pt=20, bold=True)
    return p

def add_subheading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size_pt=14, bold=True)
    return p

def add_bullet_item(doc, title="", desc=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25)
    
    r_b = p.add_run("•  ")
    set_run_font(r_b, size_pt=12, bold=True)
    
    if title:
        r_t = p.add_run(title + " ")
        set_run_font(r_t, size_pt=12, bold=True)
    if desc:
        r_d = p.add_run(desc)
        set_run_font(r_d, size_pt=12, bold=False)
    return p

def add_figure(doc, img_filename, caption_text, width_in=5.8):
    img_path = os.path.join(MEDIA_DIR, img_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(2)
        p_img.paragraph_format.keep_with_next = True
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Inches(width_in))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(6)
    
    if ":" in caption_text:
        parts = caption_text.split(":", 1)
        r_prefix = p_cap.add_run(parts[0] + ":")
        set_run_font(r_prefix, size_pt=10.5, bold=True)
        r_rest = p_cap.add_run(parts[1])
        set_run_font(r_rest, size_pt=10.5, bold=False)
    else:
        r_all = p_cap.add_run(caption_text)
        set_run_font(r_all, size_pt=10.5, bold=True)

def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

# =========================================================================
# PRELIMINARY PAGE 1: COVER / TITLE PAGE
# =========================================================================
add_p(doc, "INTERNSHIP REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_before=0, space_after=4)
add_p(doc, "ShieldJob AI – Fake Job Detection System", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_before=2, space_after=4)
add_p(doc, "Submitted in partial fulfillment of the requirements for the Summer Internship of", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_before=2, space_after=4)
add_p(doc, "Bachelor of Technology (B.Tech.)", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_before=2, space_after=6)

add_p(doc, "Submitted by", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_before=2, space_after=4)
add_p(doc, "TANUKU RAM SAI (24A95A0503)      CSE", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_before=2, space_after=6)

add_p(doc, "Under the Guidance of", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_before=2, space_after=4)
add_p(doc, "Mr. M. Srinu   Assistant Professor, Dept of CSE", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_before=2, space_after=6)

add_p(doc, "Internship Duration:\nFrom: 06/05/2026 To: 05/07/2026", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_before=2, space_after=4)
add_p(doc, "Carried out at: VedXlence Innovations Pvt. Ltd., Hyderabad", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_before=2, space_after=2)
add_p(doc, "in", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_before=2, space_after=2)
add_p(doc, "Department of Computer Science & Engineering", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_before=2, space_after=4)

emblem_path = os.path.join(MEDIA_DIR, "aditya_emblem.jpg")
if os.path.exists(emblem_path):
    p_emb = doc.add_paragraph()
    p_emb.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_emb.paragraph_format.space_before = Pt(2)
    p_emb.paragraph_format.space_after = Pt(2)
    r_emb = p_emb.add_run()
    r_emb.add_picture(emblem_path, width=Inches(1.70))

add_p(doc, "ADITYA UNIVERSITY", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_before=2, space_after=2)
add_p(doc, "Aditya Nagar, ADB Road, Surampalem, Andhra Pradesh, India", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_before=2, space_after=2)
add_p(doc, "2025–2026", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=False, space_before=2, space_after=0)

add_page_break(doc)

# =========================================================================
# PRELIMINARY PAGE 2: COLLEGE CERTIFICATE
# =========================================================================
banner_path = os.path.join(MEDIA_DIR, "aditya_banner.png")
if os.path.exists(banner_path):
    p_ban = doc.add_paragraph()
    p_ban.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ban.paragraph_format.space_before = Pt(0)
    p_ban.paragraph_format.space_after = Pt(2)
    r_ban = p_ban.add_run()
    r_ban.add_picture(banner_path, width=Inches(3.98))

add_p(doc, "Aditya Nagar, ADB Road, Surampalem, Andhra Pradesh, India", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=8.5, bold=False, space_before=0, space_after=10)
add_p(doc, "CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_before=4, space_after=10)

cert_text = (
    'This is to certify that the internship report entitled "ShieldJob AI – Fake Job Detection System" is being '
    'submitted by Tanuku Ram Sai (24A95A0503) in partial fulfillment of the requirements for the Internship '
    'in Department of Computer Science & Engineering at Aditya University. The work carried out by the student '
    'at VedXlence Innovations Pvt. Ltd., Hyderabad, between 06/05/2026 and 05/07/2026 is genuine, original, '
    'and fulfills the requirements prescribed by the institution.'
)
add_p(doc, cert_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=20, line_spacing=1.5)

sig_table = doc.add_table(rows=3, cols=2)
sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
col_widths = [Inches(3.1), Inches(3.1)]
for row in sig_table.rows:
    for i, cell in enumerate(row.cells):
        cell.width = col_widths[i]
        set_cell_margins(cell, top=40, bottom=40, left=40, right=40)

p00 = sig_table.rows[0].cells[0].paragraphs[0]
p00.alignment = WD_ALIGN_PARAGRAPH.LEFT
r00 = p00.add_run("Project Guide")
set_run_font(r00, size_pt=12, bold=True)

p01 = sig_table.rows[0].cells[1].paragraphs[0]
p01.alignment = WD_ALIGN_PARAGRAPH.LEFT
r01 = p01.add_run("Head of the Department")
set_run_font(r01, size_pt=12, bold=True)

p10 = sig_table.rows[1].cells[0].paragraphs[0]
p10.alignment = WD_ALIGN_PARAGRAPH.LEFT
r10 = p10.add_run("Signature\n\n")
set_run_font(r10, size_pt=12, bold=False)

p11 = sig_table.rows[1].cells[1].paragraphs[0]
p11.alignment = WD_ALIGN_PARAGRAPH.LEFT
r11 = p11.add_run("Signature\n\n")
set_run_font(r11, size_pt=12, bold=False)

p20 = sig_table.rows[2].cells[0].paragraphs[0]
p20.alignment = WD_ALIGN_PARAGRAPH.LEFT
r20 = p20.add_run("Name: Mr. M. Srinu\nDesignation: Assistant Professor\nDept: Computer Science & Engineering")
set_run_font(r20, size_pt=12, bold=False)

p21 = sig_table.rows[2].cells[1].paragraphs[0]
p21.alignment = WD_ALIGN_PARAGRAPH.LEFT
r21 = p21.add_run("Name: Dr. T. Sudha Rani\nDesignation: Professor & HOD\nDept: Computer Science & Engineering")
set_run_font(r21, size_pt=12, bold=False)

add_p(doc, "Date: 05/07/2026\nPlace: Surampalem", align=WD_ALIGN_PARAGRAPH.LEFT, size_pt=12, bold=False, space_before=16, space_after=0)

add_page_break(doc)

# =========================================================================
# PRELIMINARY PAGE 3: DECLARATION
# =========================================================================
if os.path.exists(banner_path):
    p_ban2 = doc.add_paragraph()
    p_ban2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_ban2.paragraph_format.space_before = Pt(0)
    p_ban2.paragraph_format.space_after = Pt(2)
    r_ban2 = p_ban2.add_run()
    r_ban2.add_picture(banner_path, width=Inches(3.98))

add_p(doc, "Aditya Nagar, ADB Road, Surampalem, Andhra Pradesh, India", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=8.5, bold=False, space_before=0, space_after=10)
add_p(doc, "DECLARATION", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_before=4, space_after=10)

decl_text1 = (
    'I hereby declare that the internship report entitled "ShieldJob AI – Fake Job Detection System" '
    'submitted to Aditya University, Surampalem, is a record of original work carried out by me during the '
    'Summer Internship period (06/05/2026 to 05/07/2026) at VedXlence Innovations Pvt. Ltd., Hyderabad, '
    'under the guidance of Mr. M. Srinu, Assistant Professor, Department of Computer Science & Engineering.'
)
add_p(doc, decl_text1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=8, line_spacing=1.5)

decl_text2 = 'I further declare that this work has not been submitted elsewhere for any degree or diploma.'
add_p(doc, decl_text2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=20, line_spacing=1.5)

add_p(doc, "Student Signature\n\n\nTANUKU RAM SAI (24A95A0503)\nDepartment of Computer Science & Engineering\nAditya University, Surampalem", 
      align=WD_ALIGN_PARAGRAPH.LEFT, size_pt=12, bold=False, space_before=8, space_after=14)

add_p(doc, "Date: 05/07/2026\nPlace: Surampalem", align=WD_ALIGN_PARAGRAPH.LEFT, size_pt=12, bold=False, space_before=8, space_after=0)

add_page_break(doc)

# =========================================================================
# PRELIMINARY PAGE 4: INTERNSHIP COMPLETION CERTIFICATE
# =========================================================================
add_p(doc, "INTERNSHIP COMPLETION CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_before=4, space_after=2)
add_p(doc, "VedXlence Innovations Pvt. Ltd., Hyderabad — AICTE Recognized Internship Program", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=11, bold=False, italic=True, space_before=2, space_after=8)

cert_img_file = os.path.join(MEDIA_DIR, "internship_certificate.png")
if os.path.exists(cert_img_file):
    t_cert = doc.add_table(rows=1, cols=1)
    t_cert.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_cert = t_cert.rows[0].cells[0]
    cell_cert.width = Inches(6.27)
    set_cell_margins(cell_cert, top=30, bottom=30, left=30, right=30)
    set_cell_background(cell_cert, "FFFFFF")
    
    p_cert_img = cell_cert.paragraphs[0]
    p_cert_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cert_img.paragraph_format.space_before = Pt(0)
    p_cert_img.paragraph_format.space_after = Pt(0)
    r_cert_img = p_cert_img.add_run()
    r_cert_img.add_picture(cert_img_file, width=Inches(6.15))
    set_table_borders(t_cert, color="888888", sz="6", val="single")

add_p(doc, "Official Certificate Verification Record — Tanuku Ram Sai (Regd. No: 24A95A0503)", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=10.5, bold=True, space_before=6, space_after=2)
add_p(doc, "Domain: Machine Learning with Python | Duration: May 6, 2026 – July 5, 2026 (2 Months)", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=10.5, bold=False, space_before=1, space_after=0)

add_page_break(doc)

# =========================================================================
# PRELIMINARY PAGE 5: ACKNOWLEDGEMENT
# =========================================================================
add_p(doc, "ACKNOWLEDGEMENT", align=WD_ALIGN_PARAGRAPH.CENTER, size_pt=12, bold=True, space_before=4, space_after=10)

ack1 = (
    "First and foremost, I would like to express my sincere gratitude to VedXlence Innovations Pvt. Ltd., "
    "Hyderabad, for providing me with the wonderful opportunity to undergo an intensive two-month AICTE-certified "
    "internship program in Machine Learning with Python. The structured training phase and rigorous real-world project "
    "implementation phase provided me with deep practical exposure to industry-standard AI development, Natural Language "
    "Processing, and full-stack deployment."
)
add_p(doc, ack1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=3, space_after=5, line_spacing=1.15)

ack2 = (
    "It is with immense pleasure that I would like to express my indebted gratitude to my project guide "
    "Mr. M. Srinu, Assistant Professor, Dept of CSE, who has guided me a lot and encouraged in every step of the "
    "project work. His valuable moral support and technical guidance throughout the project helped me to a greater extent."
)
add_p(doc, ack2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=3, space_after=5, line_spacing=1.15)

ack3 = (
    "I am profoundly grateful to Dr. T. Sudha Rani, Head of the Department of Computer Science & Engineering, "
    "for inspiring me all the way and for arranging all the facilities and resources needed for my project."
)
add_p(doc, ack3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=3, space_after=5, line_spacing=1.15)

ack4 = (
    "I wish to thank Dr. M. V. Rajesh, Associate Dean (School of Computing), Dr. A. Ramesh, Pro Vice-Chancellor "
    "(Engineering & Sciences), Dr. S. Rama Sree, Pro Vice-Chancellor (Academics), and Dr. G. Suresh, Registrar, "
    "for their encouragement and academic support during the course of my project."
)
add_p(doc, ack4, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=3, space_after=5, line_spacing=1.15)

ack5 = (
    "I would like to extend my sincere thanks to Dr. M. B. Srinivas, Vice-Chancellor, Dr. M. Sreenivasa Reddy, "
    "Deputy Pro-Chancellor, and the Management of Aditya University for their unconditional support in providing me "
    "with the best infrastructural facilities and state-of-the-art laboratories during my project."
)
add_p(doc, ack5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=3, space_after=5, line_spacing=1.15)

ack6 = (
    "Not to forget, Non-Teaching Staff and my Friends who have directly or indirectly supported me in completing "
    "this project on time."
)
add_p(doc, ack6, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=3, space_after=12, line_spacing=1.15)

add_p(doc, "TANUKU RAM SAI (24A95A0503)\nB.Tech - Computer Science & Engineering", align=WD_ALIGN_PARAGRAPH.RIGHT, size_pt=12, bold=True, space_before=8, space_after=0)

add_page_break(doc)

# =========================================================================
# PRELIMINARY PAGE 6: INDEX / TABLE OF CONTENTS
# =========================================================================
add_chapter_heading(doc, "Index")

toc_items = [
    ("1", "Abstract", "1"),
    ("2", "SDG Mapping", "2"),
    ("3", "Introduction", "3"),
    ("4", "Objectives", "4"),
    ("5", "Methodology", "5"),
    ("6", "Implementation", "7"),
    ("7", "Results and Discussion", "12"),
    ("8", "Key Learnings", "17"),
    ("9", "Conclusion and Future Scope", "18"),
    ("10", "Outcomes", "19"),
    ("11", "Publication and Product Development", "20"),
    ("12", "About the Organization", "21"),
    ("13", "References", "22")
]

toc_table = doc.add_table(rows=len(toc_items) + 1, cols=3)
toc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(toc_table, color="B0B0B0", sz="4", val="single")

toc_col_widths = [Inches(1.0), Inches(4.27), Inches(1.0)]

# Header
hdr_cells = toc_table.rows[0].cells
hdr_titles = ["S No", "Table of Content", "Page No"]
for j, cell in enumerate(hdr_cells):
    cell.width = toc_col_widths[j]
    set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
    set_cell_background(cell, "F2F2F2")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j != 1 else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(hdr_titles[j])
    set_run_font(r, size_pt=12, bold=True)

for i, item in enumerate(toc_items):
    row_cells = toc_table.rows[i + 1].cells
    for j, val in enumerate(item):
        row_cells[j].width = toc_col_widths[j]
        set_cell_margins(row_cells[j], top=50, bottom=50, left=100, right=100)
        p = row_cells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j != 1 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(val)
        set_run_font(r, size_pt=12, bold=(j == 1 and val in ["Abstract", "References", "SDG Mapping"]))

# =========================================================================
# SECTION 2: MAIN REPORT CONTENT (Numbered 1 to 22)
# =========================================================================
sec2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
sec2.page_width = Inches(8.268)
sec2.page_height = Inches(11.693)
sec2.top_margin = Inches(1.0)
sec2.bottom_margin = Inches(1.0)
sec2.left_margin = Inches(1.0)
sec2.right_margin = Inches(1.0)
sec2.header_distance = Inches(0.49)
sec2.footer_distance = Inches(0.49)
sec2.different_first_page_header_footer = False

# Unlink footer from previous section so Section 1 remains unnumbered
sec2.footer.is_linked_to_previous = False
sec2_ftr = sec2.footer.paragraphs[0]
sec2_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
sec2_ftr.text = ""
r_page = sec2_ftr.add_run()
set_run_font(r_page, size_pt=11, bold=False)
fldSimple = OxmlElement('w:fldSimple')
fldSimple.set(qn('w:instr'), 'PAGE')
r_page._r.append(fldSimple)

sectPr = sec2._sectPr
pgNumType = OxmlElement('w:pgNumType')
pgNumType.set(qn('w:start'), '1')
sectPr.append(pgNumType)

# ----------------- MAIN PAGE 1: ABSTRACT -----------------
add_chapter_heading(doc, "Abstract")

abs_text1 = (
    "Recruitment fraud on online job portals inflicts severe financial loss and identity theft on job seekers. "
    "Traditional keyword filters fail against sophisticated fraud schemes. This report presents ShieldJob AI, "
    "a multi-pillar fraud detection system combining Natural Language Processing, Machine Learning, and "
    "Generative Artificial Intelligence."
)
add_p(doc, abs_text1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

abs_text2 = (
    "Text features are extracted via TF-IDF with unigrams and bigrams across 17,880 historical job postings and classified "
    "using a Balanced Random Forest model. To overcome single-classifier limitations, Google Gemini 2.5 Flash validates "
    "contextual red flags, while BeautifulSoup scrapes company domains and search APIs audit social footprints."
)
add_p(doc, abs_text2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

abs_text3 = (
    "ShieldJob AI achieves 98.05% accuracy, 79.10% precision, 81.29% recall, and 98.49% ROC-AUC. The entire system is "
    "deployed as a full-stack Flask web application and a lightweight Manifest V3 Google Chrome Extension for real-time "
    "in-browser verification."
)
add_p(doc, abs_text3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=14, line_spacing=1.15)

add_page_break(doc)

# ----------------- MAIN PAGE 2: SDG MAPPING -----------------
add_chapter_heading(doc, "SDG Mapping")

sdg_intro = (
    "The United Nations Sustainable Development Goals (SDGs) define a universal call to action to end poverty, "
    "protect the planet, and ensure prosperity. The ShieldJob AI system directly aligns with and contributes to the "
    "following key sustainable development targets:"
)
add_p(doc, sdg_intro, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=8, line_spacing=1.15)

sdg_data = [
    ("SDG 8", "Decent Work and Economic Growth", "Target 8.8: Protect labour rights and promote safe working environments for all workers.", 
     "Protects vulnerable job seekers and students from exploitative fake employment offers, fraudulent registration fees, and employment scams, fostering a secure digital job market."),
    ("SDG 9", "Industry, Innovation and Infrastructure", "Target 9.5: Enhance scientific research, upgrade technological capabilities of industrial sectors.", 
     "Leverages modern Machine Learning (Balanced Random Forest) and Generative AI (Google Gemini) to build innovative cyber defense infrastructure against online deception."),
    ("SDG 16", "Peace, Justice and Strong Institutions", "Target 16.4: Significantly reduce illicit financial and information flows and combat cybercrime.", 
     "Deters cyber fraudsters, mitigates identity theft via stolen personal credentials, and enhances transparency and trust in online employment portals.")
]

t_sdg = doc.add_table(rows=len(sdg_data) + 1, cols=4)
t_sdg.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_sdg, color="B0B0B0", sz="4", val="single")
sdg_col_w = [Inches(0.9), Inches(1.5), Inches(1.8), Inches(2.07)]

for j, htext in enumerate(["SDG #", "Goal Title", "Relevant Target", "Project Contribution & Alignment"]):
    c = t_sdg.rows[0].cells[j]
    c.width = sdg_col_w[j]
    set_cell_margins(c, top=70, bottom=70, left=70, right=70)
    set_cell_background(c, "F2F2F2")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(htext)
    set_run_font(r, size_pt=11, bold=True)

for i, row_data in enumerate(sdg_data):
    row_cells = t_sdg.rows[i + 1].cells
    for j, val in enumerate(row_data):
        row_cells[j].width = sdg_col_w[j]
        set_cell_margins(row_cells[j], top=50, bottom=50, left=70, right=70)
        p = row_cells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(val)
        set_run_font(r, size_pt=10.5, bold=(j == 0))

add_page_break(doc)

# ----------------- MAIN PAGE 3: INTRODUCTION -----------------
add_chapter_heading(doc, "Introduction")

intro1 = (
    "The digital transformation of the employment marketplace has revolutionized modern talent acquisition. "
    "Online job boards, professional networking platforms such as LinkedIn, and automated career aggregators have "
    "drastically reduced friction in connecting job seekers with prospective employers. However, this open digital "
    "ecosystem has also fostered an exponential surge in sophisticated recruitment scams, fraudulent job advertisements, "
    "and predatory phishing campaigns targeting college graduates, entry-level professionals, and remote job seekers worldwide."
)
add_p(doc, intro1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

intro2 = (
    "Recruitment scams represent a catastrophic threat across multiple vectors: financial extortion disguised as mandatory "
    '"security deposits" or "training kits", identity theft through sensitive personal data collection (PAN, Aadhaar, bank credentials), '
    "and psychological trauma. Conventional security defenses typically rely on static keyword blacklists or isolated machine learning "
    "classifiers operating solely on raw job description text. Such single-dimensional approaches fail when fraudulent actors copy authentic "
    "company job templates verbatim, differing only in interview links, fake company domains, or unverified contact information."
)
add_p(doc, intro2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

add_subheading(doc, "The Multi-Pillar Verification Paradigm:")
intro3 = (
    "To overcome the vulnerabilities of single-method detection systems, this project presents ShieldJob AI—a comprehensive, "
    "multi-pillar fake job detection and risk assessment system developed during the internship at VedXlence Innovations Pvt. Ltd. "
    "ShieldJob AI integrates four independent analytical layers into a unified verification engine:"
)
add_p(doc, intro3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

add_bullet_item(doc, "1. Machine Learning Classifier:", "An optimized Balanced Random Forest model trained on 17,880 historical job postings using TF-IDF unigram and bigram features to detect fraudulent textual syntax.")
add_bullet_item(doc, "2. Generative AI Contextual Auditor:", "Google Gemini 2.5 Flash LLM acting as an expert fraud investigator to audit unrealistic salary claims, vague job responsibilities, and covert fee requests.")
add_bullet_item(doc, "3. Company Domain & Website Validator:", "Automated HTTP and BeautifulSoup scraping pipeline evaluating DNS records, SSL/TLS certificate validity, and corporate credibility.")
add_bullet_item(doc, "4. Social Footprint & Brand Reputation Verifier:", "Multi-engine web search verification auditing LinkedIn presence, Crunchbase registration, Glassdoor reviews, and scam alert reports.")

add_page_break(doc)

# ----------------- MAIN PAGE 4: OBJECTIVES -----------------
add_chapter_heading(doc, "Objectives")

obj_intro = "The primary engineering and research objectives achieved during the development of ShieldJob AI include:"
add_p(doc, obj_intro, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=8, line_spacing=1.15)

add_bullet_item(doc, "1. High-Accuracy Fraud Classification:", "Engineer an NLP-driven Machine Learning classification pipeline capable of exceeding 95% accuracy and 75% precision on imbalanced real-world recruitment datasets (EMSCAD).")
add_bullet_item(doc, "2. Handling Severe Class Imbalance:", "Implement balanced sub-sampling and ensemble cost-sensitive weighting to resolve the 95:5 class imbalance without overfitting or losing scam-detection sensitivity.")
add_bullet_item(doc, "3. Contextual LLM Reasoning Integration:", "Deploy Google Gemini 2.5 Flash via structured prompt engineering to extract semantic red flags and explainable risk factors invisible to pure bag-of-words models.")
add_bullet_item(doc, "4. Multi-Layer Domain & Reputation Scraping:", "Develop automated web verification services using BeautifulSoup and Search APIs to audit employer authenticity, domain age, and social footprint.")
add_bullet_item(doc, "5. Unified Composite Risk Engine:", "Design a weighted multi-factor scoring algorithm synthesizing ML probability (40%), LLM contextual score (35%), Domain reputation (15%), and Social presence (10%) into an intuitive 0–100 Safety Score.")
add_bullet_item(doc, "6. Full-Stack Web Application Deployment:", "Develop a responsive Flask Single-Page Application featuring cyber-themed UI gauges, interactive explainability cards, and PDF report export.")
add_bullet_item(doc, "7. Manifest V3 Browser Extension:", "Build a lightweight Google Chrome Extension with content scripts and background service workers to verify job postings directly on live job portals in one click.")

add_page_break(doc)

# ----------------- MAIN PAGE 5: METHODOLOGY (PART 1) -----------------
add_chapter_heading(doc, "Methodology")

meth_intro = (
    "The ShieldJob AI verification architecture combines statistical machine learning, deep contextual generative AI, "
    "and live web scraping into an end-to-end multi-pillar pipeline. This chapter details data preprocessing, feature engineering, "
    "model selection, external verification layers, and the composite scoring engine."
)
add_p(doc, meth_intro, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

add_subheading(doc, "5.1 Dataset Description and Class Distribution")
ds_text = (
    "The model was trained and evaluated on the University of the Aegean Employment Scam Aegean Dataset (EMSCAD), "
    "containing 17,880 real-world job postings. The dataset exhibits severe class imbalance, reflecting real-world conditions "
    "where fraudulent advertisements represent a critical minority:"
)
add_p(doc, ds_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

t_ds = doc.add_table(rows=4, cols=4)
t_ds.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_ds, color="B0B0B0", sz="4", val="single")
ds_col_w = [Inches(2.2), Inches(1.1), Inches(1.1), Inches(1.87)]

ds_headers = ["Category / Class", "Samples", "Percentage", "Class Weight Strategy"]
for j, htext in enumerate(ds_headers):
    c = t_ds.rows[0].cells[j]
    c.width = ds_col_w[j]
    set_cell_margins(c, top=70, bottom=70, left=70, right=70)
    set_cell_background(c, "F2F2F2")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [1, 2] else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(htext)
    set_run_font(r, size_pt=11, bold=True)

ds_rows = [
    ("Legitimate Job Postings (Class 0)", "17,014", "95.16%", "Standard Weight (Majority Class)"),
    ("Fraudulent Job Postings (Class 1)", "866", "4.84%", "Balanced Weighting (~19.6x Inverse Frequency)"),
    ("Total Training & Benchmark Corpus", "17,880", "100.00%", "5-Fold Stratified Cross-Validation")
]
for i, rdata in enumerate(ds_rows):
    rcells = t_ds.rows[i + 1].cells
    for j, val in enumerate(rdata):
        rcells[j].width = ds_col_w[j]
        set_cell_margins(rcells[j], top=50, bottom=50, left=70, right=70)
        p = rcells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in [1, 2] else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(val)
        set_run_font(r, size_pt=10.5, bold=(i == 2))

add_subheading(doc, "5.2 Data Preprocessing and Feature Engineering")
dp_text = (
    "Raw text attributes (Job Title, Company Profile, Job Description, Requirements, Benefits) are consolidated into a "
    "unified corpus. The text undergoes lowercasing, HTML stripping, punctuation filtering, and stopword removal. "
    "Numerical TF-IDF feature vectors are extracted using 10,000 maximum features across unigrams and bigrams (ngram_range=(1, 2)) "
    "with sublinear term frequency scaling (sublinear_tf=True)."
)
add_p(doc, dp_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

add_page_break(doc)

# ----------------- MAIN PAGE 6: METHODOLOGY (PART 2) -----------------
add_subheading(doc, "5.3 Multi-Pillar Verification Architecture")
arch_text = (
    "To achieve robust fraud detection resilient against sophisticated obfuscation, ShieldJob AI implements four "
    "independent verification pillars summarized below:"
)
add_p(doc, arch_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

t_arch = doc.add_table(rows=5, cols=3)
t_arch.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_arch, color="B0B0B0", sz="4", val="single")
arch_col_w = [Inches(1.5), Inches(2.1), Inches(2.67)]

for j, htext in enumerate(["Pillar Layer", "Core Technology", "Primary Inspection Function"]):
    c = t_arch.rows[0].cells[j]
    c.width = arch_col_w[j]
    set_cell_margins(c, top=70, bottom=70, left=70, right=70)
    set_cell_background(c, "F2F2F2")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(htext)
    set_run_font(r, size_pt=11, bold=True)

arch_rows = [
    ("Pillar 1: ML Classifier", "Balanced Random Forest + TF-IDF (10k n-grams)", "Analyzes structural word frequencies, syntax patterns, and suspicious keyword distributions."),
    ("Pillar 2: GenAI LLM", "Google Gemini 2.5 Flash API", "Audits contextual plausibility, unrealistic compensation, vague job duties, and covert payment demands."),
    ("Pillar 3: Domain Validator", "BeautifulSoup4 + SSL/DNS Requests", "Verifies company website reachability, SSL encryption, domain age, and career page legitimacy."),
    ("Pillar 4: Social Verifier", "Multi-Provider Web Search APIs", "Audits company reputation, LinkedIn executive profiles, Glassdoor reviews, and scam alert reports.")
]
for i, rdata in enumerate(arch_rows):
    rcells = t_arch.rows[i + 1].cells
    for j, val in enumerate(rdata):
        rcells[j].width = arch_col_w[j]
        set_cell_margins(rcells[j], top=50, bottom=50, left=70, right=70)
        p = rcells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(val)
        set_run_font(r, size_pt=10.5, bold=(j == 0))

add_subheading(doc, "5.4 Composite Risk Scoring Engine")
comp_text = (
    "The final Safety Score (0–100) is computed dynamically using a calibrated weighted fusion formula:\n"
    "• ML Fraud Probability (Weight: 40%)\n"
    "• Gemini LLM Contextual Risk Score (Weight: 35%)\n"
    "• Domain Authenticity & SSL Validity (Weight: 15%)\n"
    "• Social Footprint & Public Presence (Weight: 10%)\n\n"
    "A calibrated decision threshold of 0.24 is enforced for the ML probability to ensure high recall (~81.3%) "
    "on real scams while maintaining strong precision (~79.1%)."
)
add_p(doc, comp_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

add_page_break(doc)

# ----------------- MAIN PAGE 7: IMPLEMENTATION (PART 1) -----------------
add_chapter_heading(doc, "Implementation")

imp_intro = (
    "ShieldJob AI is engineered as a modular, production-ready full-stack software system. The codebase is partitioned "
    "into dedicated machine learning training scripts, backend Flask services, API controllers, and a Manifest V3 "
    "Google Chrome Extension. This chapter details each software module alongside implementation code architectures."
)
add_p(doc, imp_intro, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

add_subheading(doc, "6.1 Machine Learning Training Pipeline (train_model.py)")
imp1 = (
    "The training pipeline (scripts/train_model.py) loads processed datasets, builds a TF-IDF vectorization matrix, "
    "and trains a Balanced Random Forest Classifier with 5-fold cross-validation and probability threshold tuning."
)
add_p(doc, imp1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=2, space_after=4, line_spacing=1.15)
add_figure(doc, "image1.png", "Figure 6.1: Implementation of ML Training Pipeline and Stratified Cross-Validation", width_in=5.8)

add_page_break(doc)

# ----------------- MAIN PAGE 8: IMPLEMENTATION (PART 2) -----------------
add_subheading(doc, "6.2 Gemini AI Contextual Validator Service (gemini_service.py)")
imp2 = (
    "The Gemini service (services/gemini_service.py) provides centralized LLM inference using Google Gemini 2.5 Flash. "
    "It enforces strict JSON schema generation, extracts suspicious salary claims, identifies fake interview red flags, "
    "and generates human-readable risk breakdowns."
)
add_p(doc, imp2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=2, space_after=4, line_spacing=1.15)
add_figure(doc, "image2.png", "Figure 6.2: Gemini 2.5 Flash Integration with Structured Error Handling and Prompts", width_in=5.8)

add_page_break(doc)

# ----------------- MAIN PAGE 9: IMPLEMENTATION (PART 3) -----------------
add_subheading(doc, "6.3 Website & Company Domain Analyzer (website_analyzer.py)")
imp3 = (
    "The Website Analyzer (services/website_analyzer.py & services/domain_verifier.py) performs live HTTP requests, "
    "evaluates SSL certificates, checks for domain spoofing/typosquatting, and parses career pages using BeautifulSoup4."
)
add_p(doc, imp3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=2, space_after=4, line_spacing=1.15)
add_figure(doc, "image3.png", "Figure 6.3: Website Scraping and Domain Credibility Verification Implementation", width_in=5.8)

add_page_break(doc)

# ----------------- MAIN PAGE 10: IMPLEMENTATION (PART 4) -----------------
add_subheading(doc, "6.4 Social Media and Public Footprint Verifier (social_media_analyzer.py)")
imp4 = (
    "The Social Media Analyzer (services/social_media_analyzer.py) interfaces with multi-engine web search APIs "
    "(DuckDuckGo, SearXNG) to check LinkedIn executive profiles, Glassdoor ratings, and official scam alert databases."
)
add_p(doc, imp4, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=2, space_after=4, line_spacing=1.15)
add_figure(doc, "image4.png", "Figure 6.4: Social Presence Verification and Search API Integration", width_in=5.8)

add_page_break(doc)

# ----------------- MAIN PAGE 11: IMPLEMENTATION (PART 5) -----------------
add_subheading(doc, "6.5 Flask Web Application and Chrome Extension API (app.py)")
imp5 = (
    "The Flask backend (app.py) exposes the primary Single-Page Application route (/) and a lightweight JSON API "
    "endpoint (/api/analyze) consumed by both the web UI and the Manifest V3 Google Chrome Extension."
)
add_p(doc, imp5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)
add_figure(doc, "image5.png", "Figure 6.5: Flask Application Architecture and Extension API Routes", width_in=5.8)

add_page_break(doc)

# ----------------- MAIN PAGE 12: RESULTS AND DISCUSSION (PART 1) -----------------
add_chapter_heading(doc, "Results and Discussion")

res_intro = (
    "The performance of ShieldJob AI was evaluated using rigorous 5-Fold Stratified Out-of-Fold Cross-Validation "
    "across all 17,880 job postings, followed by comprehensive end-to-end testing of the live web application and Chrome extension."
)
add_p(doc, res_intro, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

add_subheading(doc, "7.1 Machine Learning Benchmark Metrics")
t_met = doc.add_table(rows=6, cols=3)
t_met.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_met, color="B0B0B0", sz="4", val="single")
met_col_w = [Inches(1.8), Inches(1.3), Inches(3.17)]

for j, htext in enumerate(["Performance Metric", "Verified Value", "Technical Interpretation & Significance"]):
    c = t_met.rows[0].cells[j]
    c.width = met_col_w[j]
    set_cell_margins(c, top=60, bottom=60, left=70, right=70)
    set_cell_background(c, "F2F2F2")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 1 else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(htext)
    set_run_font(r, size_pt=11, bold=True)

met_rows = [
    ("Accuracy", "98.05%", "Outperforms majority-class baseline with high overall predictive correctness."),
    ("Precision", "79.10%", "High reliability: ~8 out of 10 flagged jobs are confirmed fraudulent."),
    ("Recall (Sensitivity)", "81.29%", "Catches >81% of fraudulent job postings in the corpus."),
    ("F1-Score", "80.18%", "Strong harmonic mean balancing precision and recall under extreme class imbalance."),
    ("ROC-AUC Score", "98.49%", "Superb discriminative power distinguishing fraudulent vs. real postings.")
]
for i, rdata in enumerate(met_rows):
    rcells = t_met.rows[i + 1].cells
    for j, val in enumerate(rdata):
        rcells[j].width = met_col_w[j]
        set_cell_margins(rcells[j], top=40, bottom=40, left=70, right=70)
        p = rcells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j == 1 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(val)
        set_run_font(r, size_pt=10.5, bold=(j == 1))

add_subheading(doc, "7.2 Out-of-Fold Confusion Matrix Analysis")
t_cm = doc.add_table(rows=3, cols=3)
t_cm.alignment = WD_TABLE_ALIGNMENT.CENTER
set_table_borders(t_cm, color="B0B0B0", sz="4", val="single")
cm_col_w = [Inches(2.1), Inches(2.08), Inches(2.08)]

for j, htext in enumerate(["Total Samples: 17,880", "Predicted Legitimate (Real)", "Predicted Fraudulent (Fake)"]):
    c = t_cm.rows[0].cells[j]
    c.width = cm_col_w[j]
    set_cell_margins(c, top=60, bottom=60, left=70, right=70)
    set_cell_background(c, "F2F2F2")
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(htext)
    set_run_font(r, size_pt=11, bold=True)

cm_rows = [
    ("Actual Legitimate (17,014)", "True Negatives (TN): 16,828\n(Real Jobs Accurately Approved)", "False Positives (FP): 186\n(False Alarms Minimized)"),
    ("Actual Fraudulent (866)", "False Negatives (FN): 162\n(Scams Missed Minimally)", "True Positives (TP): 704\n(Scams Successfully Intercepted)")
]
for i, rdata in enumerate(cm_rows):
    rcells = t_cm.rows[i + 1].cells
    for j, val in enumerate(rdata):
        rcells[j].width = cm_col_w[j]
        set_cell_margins(rcells[j], top=40, bottom=40, left=70, right=70)
        p = rcells[j].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(val)
        set_run_font(r, size_pt=10.5, bold=False)

add_page_break(doc)

# ----------------- MAIN PAGE 13: RESULTS AND DISCUSSION (PART 2) -----------------
add_figure(doc, "image6.png", "Figure 7.1: Live Accuracy & Proof Dashboard showing 5-Fold Stratified Metrics and Confusion Matrix", width_in=5.4)

add_subheading(doc, "7.3 Working Web Application Results")
res_w = "The deployed Flask Single-Page Application provides real-time multi-layer job audits with high visual clarity:"
add_p(doc, res_w, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=2, space_after=4, line_spacing=1.15)
add_figure(doc, "image7.png", "Figure 7.2: ShieldJob AI Web Interface — Job Analysis Input Form and Options", width_in=5.4)

add_page_break(doc)

# ----------------- MAIN PAGE 14: RESULTS AND DISCUSSION (PART 3) -----------------
add_figure(doc, "image8.png", "Figure 7.3: Legitimate Job Verification Result — Emerald Status with High Safety Score", width_in=5.4)
add_figure(doc, "image9.png", "Figure 7.4: Fraudulent Job Detection Result — Rose/Red Status with Critical Risk Flags", width_in=5.4)

add_page_break(doc)

# ----------------- MAIN PAGE 15: RESULTS AND DISCUSSION (PART 4) -----------------
add_figure(doc, "image10.png", "Figure 7.5: Scan History Dashboard with Interactive Details and Re-Scan Capabilities", width_in=5.4)

add_subheading(doc, "7.4 Google Chrome Extension Verification Results")
res_c = "The Google Chrome Extension (Manifest V3) allows users to highlight or paste job descriptions directly on job portals:"
add_p(doc, res_c, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=2, space_after=4, line_spacing=1.15)
add_figure(doc, "image11.png", "Figure 7.6: Chrome Extension Popup Interface with Live Job Input Field", width_in=5.4)

add_page_break(doc)

# ----------------- MAIN PAGE 16: RESULTS AND DISCUSSION (PART 5) -----------------
add_figure(doc, "image12.png", "Figure 7.7: Chrome Extension Verification Output Displaying Safety Score and Risk Breakdown", width_in=5.8)

add_page_break(doc)

# ----------------- MAIN PAGE 17: KEY LEARNINGS -----------------
add_chapter_heading(doc, "Key Learnings")

kl_intro = (
    "The two-month internship program at VedXlence Innovations Pvt. Ltd. provided immense practical exposure "
    "across software engineering, data science, and modern AI development. Key competencies acquired include:"
)
add_p(doc, kl_intro, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=8, line_spacing=1.15)

add_bullet_item(doc, "1. Machine Learning & Imbalanced Datasets:", "Mastered techniques for handling severe class imbalances (~95:5) using Balanced Random Forest, cost-sensitive learning, stratified cross-validation, and ROC threshold calibration.")
add_bullet_item(doc, "2. Natural Language Processing (NLP):", "Gained deep practical expertise in TF-IDF feature extraction, sublinear scaling, n-gram text modeling, and tokenization pipelines.")
add_bullet_item(doc, "3. Generative AI & LLM Prompt Engineering:", "Implemented production-grade Google Gemini 2.5 Flash API calls with structured system instructions, JSON schema parsing, and deterministic error fallback.")
add_bullet_item(doc, "4. Web Scraping & DOM Parsing:", "Developed robust scraping pipelines with BeautifulSoup and requests to extract real-time web content, inspect SSL certificates, and audit domain safety.")
add_bullet_item(doc, "5. Search API Integration & Multi-Mode Fallback:", "Engineered multi-provider search services with automatic fallbacks across DuckDuckGo and SearXNG instances for high availability.")
add_bullet_item(doc, "6. Full-Stack Web Development:", "Built a modern cyber-themed Single-Page Application using Flask, HTML5, CSS3 glassmorphism, responsive dashboards, and client-side scan history caching.")
add_bullet_item(doc, "7. Browser Extension Development:", "Developed a secure Manifest V3 Chrome Extension adhering to Chrome Web Store security guidelines and background service workers.")
add_bullet_item(doc, "8. Cloud Deployment & Production Engineering:", "Successfully deployed full-stack AI applications on Render cloud infrastructure with automated environment configuration and secrets management.")

add_page_break(doc)

# ----------------- MAIN PAGE 18: CONCLUSION AND FUTURE SCOPE -----------------
add_chapter_heading(doc, "Conclusion and Future Scope")

concl_text = (
    "The ShieldJob AI project successfully demonstrates how combining classical Machine Learning with modern Generative AI "
    "and real-time web intelligence creates a powerful, explainable, and resilient defense against recruitment fraud. "
    "By achieving 98.05% accuracy, 79.10% precision, and 81.29% recall, the system provides reliable protection for job seekers "
    "and students navigating online job portals."
)
add_p(doc, concl_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=8, line_spacing=1.15)

add_subheading(doc, "9.1 Future Enhancements")
add_bullet_item(doc, "Automated DOM Job Parser:", "Enhance the Chrome Extension content scripts to automatically parse job postings from LinkedIn, Indeed, and Naukri with zero manual copy-pasting.")
add_bullet_item(doc, "Multilingual Scam Detection:", "Expand NLP tokenizers and LLM prompts to analyze regional Indian languages and identify vernacular recruitment scams.")
add_bullet_item(doc, "Company Registration & MCA Verification:", "Integrate automated checks with official corporate registries (Ministry of Corporate Affairs / ROC) to verify company registration numbers.")
add_bullet_item(doc, "Real-Time WhatsApp & Telegram Bot:", "Deploy conversational bot interfaces to allow job seekers to forward suspicious offer letters and instant messages for automated fraud analysis.")

add_page_break(doc)

# ----------------- MAIN PAGE 19: OUTCOMES -----------------
add_chapter_heading(doc, "Outcomes")

out_intro = "The concrete technical deliverables and artifacts produced as part of the internship project include:"
add_p(doc, out_intro, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=8, line_spacing=1.15)

add_bullet_item(doc, "Production Machine Learning Model:", "Serialized Balanced Random Forest classifier (models/job_fraud_model.pkl) and TF-IDF vectorizer (models/tfidf_vectorizer.pkl) with 98.05% verified accuracy.")
add_bullet_item(doc, "Full-Stack Flask Web Application:", "Responsive web app with interactive risk score gauges, explainability breakdown, scan history caching, and printable PDF audit reports.")
add_bullet_item(doc, "Google Chrome Extension (Manifest V3):", "Browser extension packaged in fake-job-detector-extension/ for direct one-click verification on live job boards.")
add_bullet_item(doc, "Live Cloud Deployment on Render:", "Hosted and accessible globally at: https://fake-job-detection-iuxn.onrender.com/#analyzer")
add_bullet_item(doc, "Open-Source GitHub Repository:", "Public codebase at: https://github.com/Coder-sai2004/Fake_Job_Detection")
add_bullet_item(doc, "Chrome Extension Repository Sub-module:", "Extension codebase at: https://github.com/Coder-sai2004/Fake_Job_Detection/tree/main/fake-job-detector-extension")

add_page_break(doc)

# ----------------- MAIN PAGE 20: PUBLICATION AND PRODUCT DEVELOPMENT -----------------
add_chapter_heading(doc, "Publication and Product Development")

add_subheading(doc, "11.1 Product Development Status")
pds_text = (
    "ShieldJob AI was engineered as a fully functional, production-ready software product rather than a theoretical prototype. "
    "Key product milestones include:"
)
add_p(doc, pds_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=6, line_spacing=1.15)

add_bullet_item(doc, "Product Name:", "ShieldJob AI – Fake Job Detection & Recruitment Risk Assessment Platform")
add_bullet_item(doc, "Product Type:", "AI/ML-Powered Cybersecurity & Employment Fraud Detection Software")
add_bullet_item(doc, "Live Product URL:", "https://fake-job-detection-iuxn.onrender.com/#analyzer")
add_bullet_item(doc, "GitHub Codebase:", "https://github.com/Coder-sai2004/Fake_Job_Detection")

add_subheading(doc, "11.2 Research Publication Status")
pub_text = (
    "Publication Status: No publication. The project was developed primarily as an industrial software product and "
    "engineering artifact during the AICTE-certified Summer Internship program at VedXlence Innovations Pvt. Ltd., "
    "prioritizing production deployment, browser extension architecture, and multi-pillar risk detection capabilities."
)
add_p(doc, pub_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=4, space_after=10, line_spacing=1.15)

add_page_break(doc)

# ----------------- MAIN PAGE 21: ABOUT THE ORGANIZATION -----------------
add_chapter_heading(doc, "About the Organization")

add_bullet_item(doc, "Company Name:", "VedXlence Innovations Pvt. Ltd.")
add_bullet_item(doc, "Company Website:", "https://www.vedxlence.in/")
add_bullet_item(doc, "Company Location:", "Hyderabad, Telangana, India")
add_bullet_item(doc, "Internship Domain:", "Machine Learning with Python")
add_bullet_item(doc, "Internship Duration:", "2 Months (May 6, 2026 – July 5, 2026)")
add_bullet_item(doc, "Program Type:", "AICTE Certified Internship & Upskilling Program")
add_bullet_item(doc, "Program Structure:", "Phase 1 – Technical Upskilling & Hands-on Training | Phase 2 – Real-World Project Implementation & Deployment")
add_bullet_item(doc, "Company Description:", "VedXlence Innovations Pvt. Ltd. is an innovative technology organization specializing in software engineering, Artificial Intelligence, Machine Learning upskilling, and enterprise digital solutions, offering AICTE-aligned technical programs for engineering students.")

add_page_break(doc)

# ----------------- MAIN PAGE 22: REFERENCES -----------------
add_chapter_heading(doc, "References")

references = [
    "[1] Pedregosa, F., et al. (2011). Scikit-learn: Machine Learning in Python. Journal of Machine Learning Research, 12, 2825-2830.",
    "[2] Google Generative AI Team. (2024). Gemini: A Family of Highly Capable Multimodal Models. Google DeepMind Technical Report.",
    "[3] Kaggle Dataset. (2020). Real or Fake: Fake JobPosting Prediction (EMSCAD). Available at: https://www.kaggle.com/datasets/shivamb/real-or-fake-fake-jobposting-prediction",
    "[4] Breiman, L. (2001). Random Forests. Machine Learning, 45(1), 5-32.",
    "[5] Ramos, J. (2003). Using TF-IDF to Determine Word Relevance in Document Queries. Proceedings of the First Instructional Conference on Machine Learning, 242-248.",
    "[6] Richardson, L. (2007). Beautiful Soup Documentation. Available at: https://www.crummy.com/software/BeautifulSoup/bs4/doc/",
    "[7] Grinberg, M. (2018). Flask Web Development: Developing Web Applications with Python (2nd ed.). O'Reilly Media.",
    "[8] Google Chrome Developers. (2023). Chrome Extensions Documentation – Manifest V3 Overview. Available at: https://developer.chrome.com/docs/extensions/mv3/intro/",
    "[9] Render Cloud Platform. (2024). Render Documentation & Deployment Guides. Available at: https://render.com/docs",
    "[10] United Nations Department of Economic and Social Affairs. (2015). The 17 Sustainable Development Goals. United Nations Guidelines."
]

for ref in references:
    add_p(doc, ref, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, space_before=3, space_after=4, line_spacing=1.15)

# Save
doc.save(OUTPUT_FILE)
print("Saved formatted document to:", OUTPUT_FILE)
doc.save(FINAL_REPLACE_FILE)
print("Updated original report at:", FINAL_REPLACE_FILE)
