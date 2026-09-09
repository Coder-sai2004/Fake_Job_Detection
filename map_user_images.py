import docx
import os
import xml.etree.ElementTree as ET

ORIGINAL_REPORT_PATH = r"C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report.docx"
doc = docx.Document(ORIGINAL_REPORT_PATH)

print("=== USER IMAGES MAPPING ===")
for i, p in enumerate(doc.paragraphs):
    xml_str = p._p.xml
    if '<a:blip' in xml_str or '<w:drawing>' in xml_str:
        # Find image target
        root = ET.fromstring(xml_str)
        target = ""
        for elem in root.iter():
            if elem.tag.endswith('blip'):
                embed = elem.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                if embed:
                    target = os.path.basename(doc.part.rels[embed].target_ref)
        # Find neighboring paragraphs (captions)
        prev_text = doc.paragraphs[i-1].text.strip() if i > 0 else ""
        curr_text = p.text.strip()
        next_text = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ""
        next2_text = doc.paragraphs[i+2].text.strip() if i+2 < len(doc.paragraphs) else ""
        print(f"P{i:03d} -> Image: {target}")
        print(f"   Prev: {prev_text[:60]}")
        print(f"   Next: {next_text[:60]}")
        print(f"   Next2: {next2_text[:60]}")
