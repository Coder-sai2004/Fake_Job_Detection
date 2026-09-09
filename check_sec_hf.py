import docx

doc = docx.Document(r"C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report_Formatted.docx")

print("Section 0 header paragraphs:", [p.text for p in doc.sections[0].header.paragraphs])
print("Section 0 footer paragraphs:", [p.text for p in doc.sections[0].footer.paragraphs])

print("Section 1 header paragraphs:", [p.text for p in doc.sections[1].header.paragraphs])
print("Section 1 footer paragraphs:", [p.text for p in doc.sections[1].footer.paragraphs])
