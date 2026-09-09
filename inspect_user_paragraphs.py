import docx

doc_u = docx.Document(r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report.docx')

print(f"Total paragraphs: {len(doc_u.paragraphs)}")
for i, p in enumerate(doc_u.paragraphs):
    t = p.text.strip()
    if t:
        print(f"P{i:03d} [{p.style.name}] (align={p.alignment}): {t[:90]}")
