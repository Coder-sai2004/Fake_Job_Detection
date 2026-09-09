import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

MEDIA_DIR = "report_assets"
OUTPUT_FILE = r"C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report_Formatted.docx"
FINAL_REPLACE_FILE = r"C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report.docx"

doc = docx.Document()

# ----------------- PAGE SETUP (A4 & 1-inch Margins) -----------------
section = doc.sections[0]
section.page_width = Inches(8.268)   # A4 Width
section.page_height = Inches(11.693) # A4 Height
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.49)
section.footer_distance = Inches(0.49)

# ----------------- HELPER STYLING FUNCTIONS -----------------
FONT_NAME = 'Times New Roman'
COLOR_BLACK = RGBColor(0, 0, 0)
COLOR_DARK = RGBColor(30, 30, 30)

def set_run_font(run, name=FONT_NAME, size_pt=12, bold=False, italic=False, color=COLOR_BLACK):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, size_pt=12, bold=False, italic=False, 
          space_before=0, space_after=6, line_spacing=1.15, keep_with_next=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.keep_with_next = keep_with_next
    if text:
        r = p.add_run(text)
        set_run_font(r, size_pt=size_pt, bold=bold, italic=italic)
    return p

def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run()
    run.add_break(docx.enum.text.WD_BREAK.PAGE)

def add_chapter_heading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size_pt=20, bold=True)
    return p

def add_subheading(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size_pt=14, bold=True)
    return p

def add_bullet_item(doc, title="", desc=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.left_indent = Inches(0.25)
    
    # Bullet character
    r_b = p.add_run("•  ")
    set_run_font(r_b, size_pt=12, bold=True)
    
    if title:
        r_t = p.add_run(title + " ")
        set_run_font(r_t, size_pt=12, bold=True)
    if desc:
        r_d = p.add_run(desc)
        set_run_font(r_d, size_pt=12, bold=False)
    return p

def add_figure(doc, img_filename, caption_text, width_in=6.0):
    img_path = os.path.join(MEDIA_DIR, img_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(3)
        p_img.paragraph_format.keep_with_next = True
        r_img = p_img.add_run()
        r_img.add_picture(img_path, width=Inches(width_in))
    
    # Caption
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(10)
    
    # Bold figure prefix
    if ":" in caption_text:
        parts = caption_text.split(":", 1)
        r_prefix = p_cap.add_run(parts[0] + ":")
        set_run_font(r_prefix, size_pt=10.5, bold=True)
        r_rest = p_cap.add_run(parts[1])
        set_run_font(r_rest, size_pt=10.5, bold=False)
    else:
        r_all = p_cap.add_run(caption_text)
        set_run_font(r_all, size_pt=10.5, bold=True)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, hex_color):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

print("Helper functions ready.")
