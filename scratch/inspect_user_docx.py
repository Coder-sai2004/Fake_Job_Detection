import docx
import os

path = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report.docx'
print("Exists:", os.path.exists(path))
print("Size:", os.path.getsize(path))

doc = docx.Document(path)
print("Paragraphs:", len(doc.paragraphs))
print("Tables:", len(doc.tables))
print("Sections:", len(doc.sections))

for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P{i:03d} [{p.style.name}]: {p.text[:100]}")
