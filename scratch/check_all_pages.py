import os
import pypdf
import docx

pdf_path = r'C:\Users\Ram Sai\Downloads\23A91A0558 Report .pdf'
reader = pypdf.PdfReader(pdf_path)

print("=== FRIEND'S PDF ALL PAGES SUMMARY ===")
for i in range(len(reader.pages)):
    txt = reader.pages[i].extract_text()
    first_lines = [line.strip() for line in txt.split('\n') if line.strip()][:3]
    print(f"Page {i+1}: {' | '.join(first_lines)}")

print("\n=== USER'S CURRENT DOCX STRUCTURE ===")
docx_path = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report.docx'
if os.path.exists(docx_path):
    doc = docx.Document(docx_path)
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    print(f"Total tables: {len(doc.tables)}")
    print(f"Total sections: {len(doc.sections)}")
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() and ("TABLE OF CONTENTS" in p.text.upper() or "INDEX" in p.text.upper() or i < 40):
            print(f"P{i:03d} [{p.style.name}]: {p.text[:90]}")
