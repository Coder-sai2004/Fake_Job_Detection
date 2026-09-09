import docx
import os
import shutil

src_temp = r'C:\Users\Ram Sai\OneDrive\Desktop\PROJECTS\Fake_Job_Detection\scratch\temp_report.docx'
doc = docx.Document(src_temp)

print("Original doc paragraphs:", len(doc.paragraphs))
print("Original doc tables:", len(doc.tables))

# Count images in original doc
img_count = sum(1 for rel in doc.part.rels.values() if "image" in rel.target_ref)
print(f"Original embedded images: {img_count}")
