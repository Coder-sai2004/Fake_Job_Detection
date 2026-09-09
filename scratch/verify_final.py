import docx
import os

path = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report_Updated.docx'
doc = docx.Document(path)
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"File size: {os.path.getsize(path)} bytes")

print("\n=== PRELIMINARY PARAGRAPHS (BEFORE TOC) ===")
toc_found = False
for i, p in enumerate(doc.paragraphs):
    if "TABLE OF CONTENTS" in p.text.upper():
        toc_found = True
        print(f"\n---> TOC STARTS AT P{i:03d}: '{p.text}' <---")
        break
    if p.text.strip():
        print(f"P{i:03d} [{p.style.name}]: {p.text[:95]}")

print("\n=== PARAGRAPHS FROM TOC ONWARD (FIRST 15) ===")
start = False
count = 0
for i, p in enumerate(doc.paragraphs):
    if "TABLE OF CONTENTS" in p.text.upper():
        start = True
    if start and p.text.strip():
        print(f"P{i:03d} [{p.style.name}]: {p.text[:95]}")
        count += 1
        if count >= 15:
            break

img_count = sum(1 for rel in doc.part.rels.values() if "image" in rel.target_ref)
print(f"\nTotal embedded images: {img_count}")
