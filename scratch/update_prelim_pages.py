import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def update_preliminary_pages():
    # Source file copied to temp
    src_temp = r'C:\Users\Ram Sai\OneDrive\Desktop\PROJECTS\Fake_Job_Detection\scratch\temp_report.docx'
    old_doc = docx.Document(src_temp)

    # Output document
    new_doc = docx.Document()

    # Match section margins (Left 1.25 in, Right 1.0 in, Top 1.0 in, Bottom 1.0 in)
    for section in new_doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.5)
        section.footer_distance = Inches(0.5)

    # Base styling
    normal_style = new_doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False, size=12, color=None, space_before=0, space_after=6, font_name="Times New Roman", line_spacing=1.15):
        p = new_doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            r = p.add_run(text)
            r.bold = bold
            r.italic = italic
            r.font.name = font_name
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        return p

    def add_bullet_p(text, bold_prefix=None, space_after=4):
        p = new_doc.add_paragraph(style='List Bullet')
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(12)
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        return p

    def add_img(img_path, width_in=4.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=6):
        if os.path.exists(img_path):
            p = new_doc.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_before = Pt(space_before)
            p.paragraph_format.space_after = Pt(space_after)
            run = p.add_run()
            run.add_picture(img_path, width=Inches(width_in))
            return p
        else:
            print(f"Warning: Image not found: {img_path}")
            return None

    scratch_dir = r'C:\Users\Ram Sai\OneDrive\Desktop\PROJECTS\Fake_Job_Detection\scratch'
    logo_path = os.path.join(scratch_dir, 'friend_page_1_img_0_Image11.jpg')
    banner2_path = os.path.join(scratch_dir, 'friend_page_2_img_0_Image19.jpg')
    banner3_path = os.path.join(scratch_dir, 'friend_page_3_img_0_Image22.jpg')
    cert_img_path = os.path.join(scratch_dir, 'user_cert_img_0_0.png')

    # =========================================================================
    # PAGE 1: TITLE / COVER PAGE (Matching 23A91A0558 Report .pdf)
    # =========================================================================
    add_p("INTERNSHIP REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_before=12, space_after=12)
    add_p("Machine Learning with Python", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, color=RGBColor(16, 44, 87), space_after=14)
    add_p("Submitted in partial fulfillment of the requirements for the Summer Internship of", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=11, space_after=2)
    add_p("Bachelor of Technology (B.Tech.)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_after=18)
    
    add_p("Submitted by", align=WD_ALIGN_PARAGRAPH.RIGHT, italic=True, size=11, space_after=2)
    add_p("Tanuku Ram Sai(24A95A0503-CSE)", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, size=12, space_after=20)

    add_p("Under the Guidance of", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=11, space_after=2)
    add_p("Dr. Jalaiah Saikam, M.Tech, Ph.D", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=12, space_after=1)
    add_p("Assistant Professor", align=WD_ALIGN_PARAGRAPH.LEFT, size=11, space_after=16)

    add_p("Internship Duration:", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=11, space_after=2)
    add_p("From: 06-05-2026 To: 05-07-2026", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=11.5, space_after=12)

    add_p("in", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=11, space_after=2)
    add_p("Department of Computer Science and Engineering", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12.5, space_after=14)

    # Aditya University Logo
    add_img(logo_path, width_in=1.8, space_before=6, space_after=10)

    add_p("ADITYA UNIVERSITY", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, color=RGBColor(16, 44, 87), space_after=1)
    add_p("(Formerly Aditya Engineering College (A))", align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, space_after=1)
    add_p("Aditya Nagar, ADB Road, Surampalem, Andhra Pradesh, India", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=1)
    add_p("2026-27", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=11, space_after=0)

    new_doc.add_page_break()

    # =========================================================================
    # PAGE 2: CERTIFICATE (Matching 23A91A0558 Report .pdf)
    # =========================================================================
    add_img(banner2_path, width_in=5.8, space_before=4, space_after=4)
    add_p("Aditya Nagar, ADB Road, Surampalem, Andhra Pradesh, India", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=4)
    add_p("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12.5, color=RGBColor(16, 44, 87), space_after=18)

    add_p("CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_after=18)

    cert_text = (
        "This is to certify that the project entitled “ShieldJob AI – Fake Job Detection System” is being submitted "
        "by Tanuku Ram Sai Roll No: 24A95A0503, in partial fulfilment of the requirements for the Internship "
        "in VedXlence Innovations Pvt. Ltd. The work carried out by the students is genuine and fulfills the requirements "
        "prescribed by the institution."
    )
    add_p(cert_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, space_after=45, line_spacing=1.3)

    add_p("Project Guide", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=12, space_after=2)
    add_p("Dr. Jalaiah Saikam, M.Tech, Ph.D", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=12, space_after=1)
    add_p("Assistant Professor", align=WD_ALIGN_PARAGRAPH.LEFT, size=11.5, space_after=25)

    add_p("Date:\nPlace: Surampalem", align=WD_ALIGN_PARAGRAPH.LEFT, size=11, space_after=0)

    new_doc.add_page_break()

    # =========================================================================
    # PAGE 3: DECLARATION (Matching 23A91A0558 Report .pdf)
    # =========================================================================
    add_img(banner3_path, width_in=5.8, space_before=4, space_after=4)
    add_p("Aditya Nagar, ADB Road, Surampalem, Andhra Pradesh, India", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=14)

    add_p("DECLARATION", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_after=18)

    decl_text = (
        "     I hereby declare that the internship report entitled \"ShieldJob AI – Fake Job Detection System\" submitted "
        "to Aditya University is a record of original work carried out by me during the summer Internship period under the guidance "
        "of Dr. Jalaiah Saikam, Assistant Professor. I further declare that this work has not been submitted elsewhere for any degree or diploma."
    )
    add_p(decl_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, space_after=40, line_spacing=1.3)

    add_p("Student Signature", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, size=11.5, space_after=4)
    add_p("Name: Tanuku Ram Sai\nRoll Number: 24A95A0503", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, size=11.5, space_after=20)
    add_p("Date:", align=WD_ALIGN_PARAGRAPH.LEFT, size=11, space_after=0)

    new_doc.add_page_break()

    # =========================================================================
    # PAGE 4: ACKNOWLEDGEMENT (Matching 23A91A0558 Report .pdf)
    # =========================================================================
    add_p("ACKNOWLEDGEMENT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_before=6, space_after=16)

    ack_p1 = (
        "It is with immense pleasure that we would like to express my indebted gratitude to my project guide "
        "Dr. Jalaiah Saikam, Department of Computer Science And Engineering, who has guided us a lot and "
        "encouraged in every step of the project work. His valuable moral support and guidance throughout the project helped me to a greater extent."
    )
    add_p(ack_p1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11.5, space_after=10, line_spacing=1.2)

    ack_p2 = (
        "We are grateful to Dr. A. Phani Sridhar HOD, for inspiring me all the way and for arranging all the "
        "facilities and resources needed for our project."
    )
    add_p(ack_p2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11.5, space_after=10, line_spacing=1.2)

    ack_p3 = (
        "We wish to thank our Dr. M.V. Rajesh, Associate Dean (School of Computing), Dr. A Ramesh, Pro Vice-Chancellor (Engineering & Sciences), "
        "Dr. S. Rama Sree, Pro Vice-Chancellor (Academics) and Dr. G. Suresh, Registrar, for their encouragement and support during the course of my project."
    )
    add_p(ack_p3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11.5, space_after=10, line_spacing=1.2)

    ack_p4 = (
        "We would like to extend my sincere thanks to Dr. M. B. Srinivas, Vice-Chancellor, Dr. M. Sreenivasa Reddy, Deputy Pro-Chancellor and Management, "
        "Aditya University for their unconditional support in providing me with the best infrastructural facilities and state-of-the-art laboratories during my project."
    )
    add_p(ack_p4, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11.5, space_after=10, line_spacing=1.2)

    ack_p5 = (
        "Not to forget, Non-Teaching Staff and our Friends who have directly or indirectly supported me in completing this project on time"
    )
    add_p(ack_p5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11.5, space_after=24, line_spacing=1.2)

    add_p("Student Name: Tanuku Ram Sai\nRegd. No: 24A95A0503", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, size=11.5, space_after=0)

    new_doc.add_page_break()

    # =========================================================================
    # PAGE 5: LEARNING OBJECTIVES / INTERNSHIP OBJECTIVES (Matching 23A91A0558 Report .pdf)
    # =========================================================================
    add_p("Learning Objectives/Internship Objectives", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=14, space_before=8, space_after=14)
    add_bullet_p("Internships are generally thought of to be reserved for college students looking to gain experience in a particular field. However, a wide array of people can benefit from Training Internships in order to receive real world experience and develop their skills.")
    add_bullet_p("An objective for this position should emphasize the skills you already possess in the area and your interest in learning more")
    add_bullet_p("Internships are utilized in a number of different career fields, including architecture, engineering, healthcare, economics, advertising and many more.")
    add_bullet_p("Some internships are used to allow individuals to perform scientific research while others are specifically designed to allow people to gain first-hand experience working.")
    add_bullet_p("Utilizing internships is a great way to build your resume and develop skills that can be emphasized in your resume for future jobs. When you are applying for a Training Internship, make sure to highlight any special skills or talents that can make you stand apart from the rest of the applicants so that you have an improved chance of landing the position.")

    new_doc.add_page_break()

    # =========================================================================
    # PAGE 6: INTERNSHIP COMPLETION CERTIFICATE (Matching 23A91A0558 Report .pdf)
    # =========================================================================
    add_p("INTERNSHIP COMPLETION CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_before=10, space_after=14)
    if os.path.exists(cert_img_path):
        add_img(cert_img_path, width_in=5.8, space_before=10, space_after=10)
    else:
        add_p("[ Official Internship Certificate Copy Attached — Tanuku Ram Sai (24A95A0503) ]\nIssued by: VedXlence Innovations Pvt. Ltd., Hyderabad\nAICTE Certified Upskilling & Project Implementation Program", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, italic=True, size=11)

    new_doc.add_page_break()

    # =========================================================================
    # PRESERVE EVERYTHING FROM TABLE OF CONTENTS ONWARD
    # =========================================================================
    # Find the TOC paragraph in old_doc
    toc_idx = -1
    for i, p in enumerate(old_doc.paragraphs):
        if "TABLE OF CONTENTS" in p.text.upper():
            toc_idx = i
            break

    print(f"Copying content from TOC onward (starting at old_doc paragraph P{toc_idx})...")
    
    # We will copy all elements in body order from TOC onward
    # Let's inspect the body elements in old_doc to ensure tables and images are copied exactly
    body = old_doc._body._element
    copying = False
    
    # Let's copy using python-docx XML elements or paragraph/table replication
    for child in body:
        tag = child.tag.split('}')[-1]
        if tag == 'p':
            p_text = "".join(child.itertext()).strip()
            if "TABLE OF CONTENTS" in p_text.upper():
                copying = True
            if copying:
                new_doc._body._element.append(child)
        elif tag == 'tbl':
            if copying:
                new_doc._body._element.append(child)
        elif tag == 'sectPr':
            pass

    out_file1 = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report_Updated.docx'
    out_file2 = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\24A95A0503_ShieldJob_AI_Internship_Report.docx'
    out_file_orig = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report.docx'

    new_doc.save(out_file1)
    print(f"Saved updated report to: {out_file1}")
    
    new_doc.save(out_file2)
    print(f"Saved updated report to: {out_file2}")

    try:
        new_doc.save(out_file_orig)
        print(f"Successfully overwritten: {out_file_orig}")
    except Exception as e:
        print(f"Could not overwrite {out_file_orig} directly (file is open in Word): {e}")

if __name__ == "__main__":
    update_preliminary_pages()
