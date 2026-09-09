import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    COLOR_BG = RGBColor(248, 250, 252)        # Light Slate #F8FAFC
    COLOR_PRIMARY = RGBColor(15, 23, 42)      # Deep Navy #0F172A
    COLOR_SECONDARY = RGBColor(30, 41, 59)    # Slate #1E293B
    COLOR_ACCENT = RGBColor(14, 165, 233)     # Cyan/Blue #0EA5E9
    COLOR_ACCENT_DARK = RGBColor(37, 99, 235) # Royal Blue #2563EB
    COLOR_MUTED = RGBColor(100, 116, 139)     # Slate 500 #64748B
    COLOR_CARD_BG = RGBColor(255, 255, 255)   # White
    COLOR_CARD_BORDER = RGBColor(226, 232, 240) # Slate 200 #E2E8F0
    COLOR_SUCCESS = RGBColor(16, 185, 129)    # Emerald 500 #10B981

    def set_slide_background(slide, color=COLOR_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="INTERNSHIP REVIEW"):
        # Top banner accent line
        accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(0.12), Inches(0.85))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = COLOR_ACCENT
        accent_bar.line.fill.background()

        # Category & Title text
        tx_box = slide.shapes.add_textbox(Inches(1.05), Inches(0.42), Inches(11.4), Inches(0.95))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = category_text.upper()
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_ACCENT_DARK
        p1.font.name = "Segoe UI"

        p2 = tf.add_paragraph()
        p2.text = title_text
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = COLOR_PRIMARY
        p2.font.name = "Segoe UI"

    def add_card(slide, left, top, width, height, title="", items=None, accent_top=False):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_CARD_BG
        card.line.color.rgb = COLOR_CARD_BORDER
        card.line.width = Pt(1.5)

        if accent_top:
            acc = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + Inches(0.1), top, width - Inches(0.2), Inches(0.06))
            acc.fill.solid()
            acc.fill.fore_color.rgb = COLOR_ACCENT
            acc.line.fill.background()

        tx_box = slide.shapes.add_textbox(left + Inches(0.3), top + Inches(0.25), width - Inches(0.6), height - Inches(0.5))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        if title:
            p_title = tf.paragraphs[0]
            p_title.text = title
            p_title.font.size = Pt(18)
            p_title.font.bold = True
            p_title.font.color.rgb = COLOR_PRIMARY
            p_title.font.name = "Segoe UI"
            p_title.space_after = Pt(12)

        if items:
            for idx, item in enumerate(items):
                p = tf.add_paragraph() if (title or idx > 0) else tf.paragraphs[0]
                p.text = f"•   {item}"
                p.font.size = Pt(14)
                p.font.color.rgb = COLOR_SECONDARY
                p.font.name = "Segoe UI"
                p.space_after = Pt(8)
                p.line_spacing = 1.15

        return card

    # ==================== SLIDE 1: INTERNSHIP OVERVIEW ====================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1)
    
    # Title Hero Card
    hero = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.73), Inches(2.2))
    hero.fill.solid()
    hero.fill.fore_color.rgb = COLOR_PRIMARY
    hero.line.fill.background()

    htf = hero.text_frame
    htf.word_wrap = True
    htf.margin_left = Inches(0.5)
    htf.margin_top = Inches(0.35)
    
    hp0 = htf.paragraphs[0]
    hp0.text = "INTERNSHIP REVIEW PRESENTATION"
    hp0.font.size = Pt(13)
    hp0.font.bold = True
    hp0.font.color.rgb = COLOR_ACCENT
    hp0.font.name = "Segoe UI"

    hp1 = htf.add_paragraph()
    hp1.text = "ShieldJob AI: Multi-Modal Fake Job Detection System"
    hp1.font.size = Pt(28)
    hp1.font.bold = True
    hp1.font.color.rgb = RGBColor(255, 255, 255)
    hp1.font.name = "Segoe UI"
    hp1.space_after = Pt(4)

    hp2 = htf.add_paragraph()
    hp2.text = "AICTE Certified Internship & Upskilling Program | Machine Learning with Python"
    hp2.font.size = Pt(14)
    hp2.font.color.rgb = RGBColor(203, 213, 225)
    hp2.font.name = "Segoe UI"

    # 4 Overview Cards
    cards_data = [
        ("Company", ["VedXlence Innovations Pvt. Ltd.", "Hyderabad, Telangana, India"]),
        ("Domain", ["Machine Learning with Python", "Natural Language Processing"]),
        ("Duration & Dates", ["2 Months Duration", "May 6, 2026 – July 5, 2026"]),
        ("Mode & Program", ["Offline Mode", "AICTE Certified Program"])
    ]
    card_w = Inches(2.76)
    card_h = Inches(3.6)
    for i, (ctitle, citems) in enumerate(cards_data):
        c_left = Inches(0.8) + i * Inches(2.98)
        add_card(slide1, c_left, Inches(3.25), card_w, card_h, title=ctitle, items=citems, accent_top=True)

    # ==================== SLIDE 2: ORGANIZATION ====================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2)
    add_header(slide2, "Organization Overview", "About the Company")

    add_card(slide2, Inches(0.8), Inches(1.6), Inches(11.73), Inches(5.2), title="VedXlence Innovations Pvt. Ltd.", items=[
        "Organization: VedXlence Innovations Pvt. Ltd., based in Hyderabad, Telangana, India.",
        "Program Type: AICTE Certified Internship & Upskilling Program designed for engineering graduates.",
        "Core Focus: Structured technical training followed by industry-aligned real-world project implementation.",
        "Skill Enablement: Bridging academic machine learning concepts with practical, production-ready software solutions.",
        "Official Portal: https://www.vedxlence.in/"
    ], accent_top=True)

    # ==================== SLIDE 3: WORK & RESPONSIBILITIES ====================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3)
    add_header(slide3, "Internship Work & Responsibilities", "Execution & Pipeline")

    add_card(slide3, Inches(0.8), Inches(1.6), Inches(5.72), Inches(5.2), title="ML Pipeline & Data Engineering", items=[
        "Data Preprocessing & Analysis: Cleaned and structured 17,880 job records from the EMSCAD dataset.",
        "NLP & Feature Engineering: Applied TF-IDF vectorization to convert raw job texts into informative numerical feature vectors.",
        "ML Model Training: Trained and optimized a Balanced Random Forest classifier addressing dataset imbalance.",
        "Testing & Metrics: Evaluated models using accuracy, precision, recall, F1-score, and ROC-AUC curves."
    ], accent_top=True)

    add_card(slide3, Inches(6.8), Inches(1.6), Inches(5.72), Inches(5.2), title="AI, Web & Deployment Integration", items=[
        "Generative AI Integration: Integrated Gemini 2.5 Flash for semantic contextual analysis and red-flag extraction.",
        "Automated Verification: Implemented automated website status and public social footprint checking.",
        "Application Development: Built an interactive Flask web application and a Chrome Extension (Manifest V3).",
        "Cloud Deployment: Handled error resilience and deployed the full application to Render cloud platform."
    ], accent_top=True)

    # ==================== SLIDE 4: SKILLS & TECHNOLOGIES LEARNED ====================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4)
    add_header(slide4, "Skills & Technologies Learned", "Technical Competencies")

    tech_cards = [
        ("Python & Flask", ["Core Python 3", "Flask Framework", "Gunicorn Server", "RESTful API Design"]),
        ("Machine Learning & NLP", ["Scikit-learn Library", "TF-IDF Vectorization", "Balanced Random Forest", "Model Evaluation Metrics"]),
        ("Generative AI & Web", ["Gemini 2.5 Flash API", "Prompt Engineering", "BeautifulSoup Scraping", "Social/Website Verifiers"]),
        ("Frontend & Deployment", ["HTML5, CSS3, JavaScript", "Chrome Extension (V3)", "Git & GitHub Workflow", "Render Cloud Hosting"])
    ]
    for i, (ttitle, titems) in enumerate(tech_cards):
        c_left = Inches(0.8) + i * Inches(2.98)
        add_card(slide4, c_left, Inches(1.6), Inches(2.76), Inches(5.2), title=ttitle, items=titems, accent_top=True)

    # ==================== SLIDE 5: PROJECT INTRODUCTION ====================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5)
    add_header(slide5, "Project Introduction: ShieldJob AI", "Capstone Project")

    add_card(slide5, Inches(0.8), Inches(1.6), Inches(11.73), Inches(5.2), title="ShieldJob AI – Multi-Modal Fake Job Detection System", items=[
        "The Problem: Online recruitment scams and fraudulent job postings have surged, putting job seekers at serious financial and privacy risk.",
        "The Objective: Automatically identify whether a job posting is legitimate, suspicious, or potentially fraudulent with high precision.",
        "The Multi-Modal Approach: Unlike traditional systems relying solely on text keywords, ShieldJob AI combines Machine Learning, Generative AI, domain verification, and digital footprint analysis.",
        "The Solution: Provides an explainable, multi-layer verification score (0–100) accessible via Web Application and Chrome Extension."
    ], accent_top=True)

    # ==================== SLIDE 6: PROJECT WORKFLOW ====================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6)
    add_header(slide6, "System Workflow & Verification Weights", "Architecture")

    workflow_stages = [
        ("1. Input Submission", ["User inputs Job Posting text or checks live job page via Chrome Extension."]),
        ("2. Multi-Modal Checks", [
            "ML Text Classifier (30%)",
            "Gemini AI Analysis (30%)",
            "Website Verification (20%)",
            "Social Footprint Check (20%)"
        ]),
        ("3. Score Computation", ["Calculates unified Safety Score (0–100) based on weighted confidence levels."]),
        ("4. Final Result", ["Displays clear verdict (Legitimate / Suspicious / Fraudulent) with explainable red flags."])
    ]
    for i, (stitle, sitems) in enumerate(workflow_stages):
        c_left = Inches(0.8) + i * Inches(2.98)
        add_card(slide6, c_left, Inches(1.6), Inches(2.76), Inches(5.2), title=stitle, items=sitems, accent_top=True)

    # ==================== SLIDE 7: PROJECT FEATURES ====================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7)
    add_header(slide7, "Key Features & Capabilities", "System Capabilities")

    add_card(slide7, Inches(0.8), Inches(1.6), Inches(5.72), Inches(5.2), title="Detection & Verification Core", items=[
        "Job Scam Detection: Multi-layered verification catching deceptive job postings.",
        "AI-Based Content Analysis: Gemini AI detects subtle fraudulent phrasing, unrealistic salaries, and suspicious contacts.",
        "Company Website Verification: Confirms domain accessibility, validity, and authentic corporate presence.",
        "Public Footprint Verification: Cross-references public digital and social indicators."
    ], accent_top=True)

    add_card(slide7, Inches(6.8), Inches(1.6), Inches(5.72), Inches(5.2), title="User Experience & Access", items=[
        "0–100 Safety Score: Single, intuitive trust index for rapid decision making.",
        "Explainable Red Flags: Clear, human-readable explanations of all detected anomalies.",
        "Flask Web Application: Full portal with comprehensive breakdown and report view.",
        "Chrome Extension (V3): Real-time analysis directly on live job portals without leaving the browser tab."
    ], accent_top=True)

    # ==================== SLIDE 8: RESULTS / OUTCOME ====================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8)
    add_header(slide8, "Experimental Results & Model Performance", "Performance & Metrics")

    # Big Accuracy Highlight Card
    acc_card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(4.5), Inches(5.2))
    acc_card.fill.solid()
    acc_card.fill.fore_color.rgb = COLOR_PRIMARY
    acc_card.line.fill.background()

    atf = acc_card.text_frame
    atf.word_wrap = True
    atf.margin_left = Inches(0.4)
    atf.margin_top = Inches(0.4)

    ap0 = atf.paragraphs[0]
    ap0.text = "PRIMARY RESULT"
    ap0.font.size = Pt(12)
    ap0.font.bold = True
    ap0.font.color.rgb = COLOR_ACCENT
    ap0.font.name = "Segoe UI"

    ap1 = atf.add_paragraph()
    ap1.text = "98.05%"
    ap1.font.size = Pt(48)
    ap1.font.bold = True
    ap1.font.color.rgb = COLOR_SUCCESS
    ap1.font.name = "Segoe UI"
    ap1.space_after = Pt(8)

    ap2 = atf.add_paragraph()
    ap2.text = "Overall Model Accuracy achieved by Balanced Random Forest Classifier on the EMSCAD benchmark dataset."
    ap2.font.size = Pt(14)
    ap2.font.color.rgb = RGBColor(226, 232, 240)
    ap2.font.name = "Segoe UI"
    ap2.space_after = Pt(14)

    ap3 = atf.add_paragraph()
    ap3.text = "• Dataset: EMSCAD (17,880 Postings)\n• Cloud Status: Live on Render\n• Source: GitHub Repository"
    ap3.font.size = Pt(13)
    ap3.font.color.rgb = RGBColor(148, 163, 184)
    ap3.font.name = "Segoe UI"

    # Evaluation Metrics Card
    add_card(slide8, Inches(5.6), Inches(1.6), Inches(6.93), Inches(5.2), title="Detailed Model Evaluation Metrics", items=[
        "Dataset Benchmark: 17,880 real-world job postings from the EMSCAD dataset.",
        "Model Architecture: Balanced Random Forest addressing class imbalance.",
        "Model Accuracy: 98.05% overall classification accuracy.",
        "Precision: 79.10% (Minimizing false fraud accusations).",
        "Recall: 81.29% (High capture rate of fraudulent job postings).",
        "F1-Score: 80.18% (Balanced harmonic mean between precision and recall).",
        "ROC-AUC: 98.49% (Exceptional discriminatory power between classes).",
        "Deployment: Successfully hosted and functional on Render cloud platform."
    ], accent_top=True)

    # ==================== SLIDE 9: KEY LEARNINGS ====================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide9)
    add_header(slide9, "Key Learnings & Takeaways", "Growth & Experience")

    add_card(slide9, Inches(0.8), Inches(1.6), Inches(5.72), Inches(5.2), title="Technical Competencies", items=[
        "Practical ML Pipeline: Data cleaning, feature extraction (TF-IDF), and ensemble model tuning.",
        "Generative AI Integration: Real-world prompt design and structured API integration using Gemini AI.",
        "Web Scraping & APIs: Automated extraction, domain validation, and external service integration.",
        "Browser Extension Architecture: Built interactive Chrome Extension communicating with Flask backend."
    ], accent_top=True)

    add_card(slide9, Inches(6.8), Inches(1.6), Inches(5.72), Inches(5.2), title="Software Engineering Experience", items=[
        "End-to-End System Development: Integrated ML models into a fully functioning multi-tier web application.",
        "Error Resilience & Testing: Handled network timeouts, missing fields, and edge cases gracefully.",
        "Cloud Deployment: Packaged and deployed production service with Gunicorn and Render.",
        "Project Management: Maintained version control, modular code architecture, and Git workflows."
    ], accent_top=True)

    # ==================== SLIDE 10: DEMO TRANSITION ====================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide10)

    demo_card = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), Inches(10.33), Inches(4.5))
    demo_card.fill.solid()
    demo_card.fill.fore_color.rgb = COLOR_PRIMARY
    demo_card.line.fill.background()

    dtf = demo_card.text_frame
    dtf.word_wrap = True
    dtf.margin_left = Inches(0.8)
    dtf.margin_top = Inches(0.8)

    dp0 = dtf.paragraphs[0]
    dp0.text = "LIVE DEMONSTRATION"
    dp0.font.size = Pt(14)
    dp0.font.bold = True
    dp0.font.color.rgb = COLOR_ACCENT
    dp0.font.name = "Segoe UI"

    dp1 = dtf.add_paragraph()
    dp1.text = "ShieldJob AI System Demo"
    dp1.font.size = Pt(36)
    dp1.font.bold = True
    dp1.font.color.rgb = RGBColor(255, 255, 255)
    dp1.font.name = "Segoe UI"
    dp1.space_after = Pt(16)

    dp2 = dtf.add_paragraph()
    dp2.text = "Demonstrating the live Flask Web Application & Chrome Extension with real-time job scam detection."
    dp2.font.size = Pt(18)
    dp2.font.color.rgb = RGBColor(203, 213, 225)
    dp2.font.name = "Segoe UI"

    # ==================== SPEAKING SCRIPT SLIDES (SEPARATE AT THE END) ====================
    
    # Script Slide 1 (Slides 1 to 3)
    ss1 = prs.slides.add_slide(blank_layout)
    set_slide_background(ss1)
    add_header(ss1, "Speaking Script: Slides 1 to 3", "Speaker Script (For Presentation Delivery)")
    add_card(ss1, Inches(0.8), Inches(1.6), Inches(11.73), Inches(5.2), title="Slide-by-Slide Delivery Script", items=[
        "[Slide 1 — Overview | ~35 sec]: \"Respected panel members and faculty, good morning. Today, I am presenting my internship review on the work I completed at VedXlence Innovations Pvt. Ltd., located in Hyderabad. I completed a two-month offline internship in Machine Learning with Python, from May 6, 2026, to July 5, 2026, as part of an AICTE-certified internship and upskilling program.\"",
        "[Slide 2 — Organization | ~30 sec]: \"VedXlence Innovations Pvt. Ltd. provided an AICTE-certified internship and upskilling program. The organization focused on structured technical training followed by practical, real-world project implementation, allowing us to bridge academic learning with industry development practices.\"",
        "[Slide 3 — Work & Responsibilities | ~45 sec]: \"During this internship, my responsibilities covered the complete pipeline: data preprocessing, NLP with TF-IDF, training ML classification models, integrating Gemini AI, building website/social verification checks, developing a Flask app and Chrome extension, and deploying to Render.\""
    ], accent_top=True)

    # Script Slide 2 (Slides 4 to 6)
    ss2 = prs.slides.add_slide(blank_layout)
    set_slide_background(ss2)
    add_header(ss2, "Speaking Script: Slides 4 to 6", "Speaker Script (For Presentation Delivery)")
    add_card(ss2, Inches(0.8), Inches(1.6), Inches(11.73), Inches(5.2), title="Slide-by-Slide Delivery Script", items=[
        "[Slide 4 — Skills Learned | ~40 sec]: \"I gained hands-on competencies in Python, Flask, Scikit-learn, and NLP. I also learned Generative AI integration with Gemini 2.5 Flash, web scraping with BeautifulSoup, JavaScript for Chrome extensions, Git version control, and cloud hosting on Render.\"",
        "[Slide 5 — Project Introduction | ~45 sec]: \"As the capstone project, I developed ShieldJob AI – a Multi-Modal Fake Job Detection System. With online recruitment scams rising, ShieldJob AI automatically evaluates whether a job posting is legitimate, suspicious, or fraudulent by combining text classification, GenAI, and company footprint verification.\"",
        "[Slide 6 — Project Workflow | ~45 sec]: \"The workflow combines four concurrent checks: ML text analysis (30%), Gemini AI analysis (30%), company website check (20%), and social footprint check (20%). These compute a unified 0–100 Safety Score to produce the final classification and red flags.\""
    ], accent_top=True)

    # Script Slide 3 (Slides 7 to 10)
    ss3 = prs.slides.add_slide(blank_layout)
    set_slide_background(ss3)
    add_header(ss3, "Speaking Script: Slides 7 to 10", "Speaker Script (For Presentation Delivery)")
    add_card(ss3, Inches(0.8), Inches(1.6), Inches(11.73), Inches(5.2), title="Slide-by-Slide Delivery Script", items=[
        "[Slide 7 — Key Features | ~40 sec]: \"Key features include comprehensive scam detection, AI-based content analysis, domain verification, a 0–100 safety score, explainable red flags, and dual access via a Flask Web App and a live Chrome Extension.\"",
        "[Slide 8 — Results | ~40 sec]: \"Trained on 17,880 postings from the EMSCAD dataset, our Balanced Random Forest model achieved a high 98.05% overall accuracy, 98.49% ROC-AUC, 79.10% precision, and 81.29% recall, successfully deployed on Render.\"",
        "[Slide 9 — Key Learnings | ~35 sec]: \"This internship gave me comprehensive experience in ML workflows, GenAI integration, full-stack web and extension development, and production cloud deployment.\"",
        "[Slide 10 — Demo Transition | ~15 sec]: \"I will now transition to the live demonstration of the ShieldJob AI web application and Chrome extension to show the system working in real time. Thank you.\""
    ], accent_top=True)

    # Save to both target locations
    target_path = r"C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\Fake_Job_Detection_Internship_Presentation.pptx"
    local_path = r"c:\Users\Ram Sai\OneDrive\Desktop\PROJECTS\Fake_Job_Detection\Fake_Job_Detection_Internship_Presentation.pptx"
    
    os.makedirs(os.path.dirname(target_path), exist_ok=True)
    prs.save(target_path)
    prs.save(local_path)
    print(f"Successfully generated presentation at:\n1. {target_path}\n2. {local_path}")

if __name__ == "__main__":
    create_presentation()
