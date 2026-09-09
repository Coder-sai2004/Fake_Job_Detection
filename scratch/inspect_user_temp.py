import docx
import os

path = r'C:\Users\Ram Sai\OneDrive\Desktop\PROJECTS\Fake_Job_Detection\scratch\temp_report.docx'
doc = docx.Document(path)
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

print("\n--- ALL PARAGRAPHS IN USER'S DOCX ---")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt:
        print(f"P{i:03d} [{p.style.name}]: {txt[:100]}")

print("\n--- ALL TABLES IN USER'S DOCX ---")
for i, t in enumerate(doc.tables):
    print(f"\nTable {i}: {len(t.rows)} rows x {len(t.columns)} cols")
    for r_idx, r in enumerate(t.rows):
        row_txt = [c.text.strip().replace('\n', ' ') for c in r.cells]
        print(f"  R{r_idx}: {' | '.join(row_txt[:4])}")
