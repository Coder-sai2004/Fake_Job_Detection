import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def perform_in_place_update():
    src_temp = r'C:\Users\Ram Sai\OneDrive\Desktop\PROJECTS\Fake_Job_Detection\scratch\temp_report.docx'
    doc = docx.Document(src_temp)

    scratch_dir = r'C:\Users\Ram Sai\OneDrive\Desktop\PROJECTS\Fake_Job_Detection\scratch'
    logo_path = os.path.join(scratch_dir, 'friend_page_1_img_0_Image11.jpg')
    banner2_path = os.path.join(scratch_dir, 'friend_page_2_img_0_Image19.jpg')
    banner3_path = os.path.join(scratch_dir, 'friend_page_3_img_0_Image22.jpg')
    cert_img_path = os.path.join(scratch_dir, 'user_cert_img_0_0.png')

    # Find the TOC element in document body
    toc_p = None
    for p in doc.paragraphs:
        if "TABLE OF CONTENTS" in p.text.upper():
            toc_p = p
            break

    if toc_p is None:
        print("Error: Could not find TABLE OF CONTENTS paragraph!")
        return

    # Remove all paragraphs and tables that appear BEFORE the TOC paragraph element
    toc_elem = toc_p._element
    parent = toc_elem.getparent()

    elements_to_remove = []
    for child in parent:
        if child == toc_elem:
            break
        tag = child.tag.split('}')[-1]
        if tag in ['p', 'tbl']:
            elements_to_remove.append(child)

    print(f"Removing {len(elements_to_remove)} preliminary elements before TOC...")
    for elem in elements_to_remove:
        parent.remove(elem)

    # Helper to insert a paragraph before toc_elem
    def insert_p_before(text="", align=WD_ALIGN_PARAGRAPH.LEFT, bold=False, italic=False, size=12, color=None, space_before=0, space_after=6, font_name="Times New Roman", line_spacing=1.15):
        new_p = doc.add_paragraph() # creates paragraph at end
        new_p.alignment = align
        new_p.paragraph_format.space_before = Pt(space_before)
        new_p.paragraph_format.space_after = Pt(space_after)
        new_p.paragraph_format.line_spacing = line_spacing
        if text:
            r = new_p.add_run(text)
            r.bold = bold
            r.italic = italic
            r.font.name = font_name
            r.font.size = Pt(size)
            if color:
                r.font.color.rgb = color
        # move element before toc_elem
        toc_elem.addprevious(new_p._element)
        return new_p

    def insert_bullet_before(text, space_after=4):
        new_p = doc.add_paragraph(style='List Bullet')
        new_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        new_p.paragraph_format.space_after = Pt(space_after)
        new_p.paragraph_format.line_spacing = 1.15
        r = new_p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        toc_elem.addprevious(new_p._element)
        return new_p

    def insert_img_before(img_path, width_in=4.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=6):
        if os.path.exists(img_path):
            new_p = doc.add_paragraph()
            new_p.alignment = align
            new_p.paragraph_format.space_before = Pt(space_before)
            new_p.paragraph_format.space_after = Pt(space_after)
            run = new_p.add_run()
            run.add_picture(img_path, width=Inches(width_in))
            toc_elem.addprevious(new_p._element)
            return new_p
        return None

    def insert_page_break_before():
        new_p = doc.add_paragraph()
        new_p.add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
        toc_elem.addprevious(new_p._element)

    # =========================================================================
    # 1. TITLE / COVER PAGE
    # =========================================================================
    insert_p_before("INTERNSHIP REPORT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_before=10, space_after=10)
    insert_p_before("Machine Learning with Python", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=16, color=RGBColor(16, 44, 87), space_after=12)
    insert_p_before("Submitted in partial fulfillment of the requirements for the Summer Internship of", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=11, space_after=2)
    insert_p_before("Bachelor of Technology (B.Tech.)", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=13, space_after=16)
    
    insert_p_before("Submitted by", align=WD_ALIGN_PARAGRAPH.RIGHT, italic=True, size=11, space_after=2)
    insert_p_before("Tanuku Ram Sai(24A95A0503-CSE)", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, size=12, space_after=18)

    insert_p_before("Under the Guidance of", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=11, space_after=2)
    insert_p_before("Dr. Jalaiah Saikam, M.Tech, Ph.D", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=12, space_after=1)
    insert_p_before("Assistant Professor", align=WD_ALIGN_PARAGRAPH.LEFT, size=11, space_after=14)

    insert_p_before("Internship Duration:", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=11, space_after=2)
    insert_p_before("From: 06-05-2026 To: 05-07-2026", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=11.5, space_after=10)

    insert_p_before("in", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, size=11, space_after=2)
    insert_p_before("Department of Computer Science and Engineering", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12.5, space_after=12)

    insert_img_before(logo_path, width_in=1.8, space_before=4, space_after=8)

    insert_p_before("ADITYA UNIVERSITY", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=14, color=RGBColor(16, 44, 87), space_after=1)
    insert_p_before("(Formerly Aditya Engineering College (A))", align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5, space_after=1)
    insert_p_before("Aditya Nagar, ADB Road, Surampalem, Andhra Pradesh, India", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=1)
    insert_p_before("2026-27", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=11, space_after=0)

    insert_page_break_before()

    # =========================================================================
    # 2. CERTIFICATE
    # =========================================================================
    insert_img_before(banner2_path, width_in=5.8, space_before=4, space_after=4)
    insert_p_before("Aditya Nagar, ADB Road, Surampalem, Andhra Pradesh, India", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=4)
    insert_p_before("DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=12.5, color=RGBColor(16, 44, 87), space_after=16)

    insert_p_before("CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_after=16)

    cert_text = (
        "This is to certify that the project entitled “ShieldJob AI – Fake Job Detection System” is being submitted "
        "by Tanuku Ram Sai Roll No: 24A95A0503, in partial fulfilment of the requirements for the Internship "
        "in VedXlence Innovations Pvt. Ltd. The work carried out by the students is genuine and fulfills the requirements "
        "prescribed by the institution."
    )
    insert_p_before(cert_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, space_after=40, line_spacing=1.3)

    insert_p_before("Project Guide", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=12, space_after=2)
    insert_p_before("Dr. Jalaiah Saikam, M.Tech, Ph.D", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=12, space_after=1)
    insert_p_before("Assistant Professor", align=WD_ALIGN_PARAGRAPH.LEFT, size=11.5, space_after=25)

    insert_p_before("Date:\nPlace: Surampalem", align=WD_ALIGN_PARAGRAPH.LEFT, size=11, space_after=0)

    insert_page_break_before()

    # =========================================================================
    # 3. DECLARATION
    # =========================================================================
    insert_img_before(banner3_path, width_in=5.8, space_before=4, space_after=4)
    insert_p_before("Aditya Nagar, ADB Road, Surampalem, Andhra Pradesh, India", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, space_after=12)

    insert_p_before("DECLARATION", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_after=16)

    decl_text = (
        "     I hereby declare that the internship report entitled \"ShieldJob AI – Fake Job Detection System\" submitted "
        "to Aditya University is a record of original work carried out by me during the summer Internship period under the guidance "
        "of Dr. Jalaiah Saikam, Assistant Professor. I further declare that this work has not been submitted elsewhere for any degree or diploma."
    )
    insert_p_before(decl_text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=12, space_after=35, line_spacing=1.3)

    insert_p_before("Student Signature", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, size=11.5, space_after=4)
    insert_p_before("Name: Tanuku Ram Sai\nRoll Number: 24A95A0503", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, size=11.5, space_after=20)
    insert_p_before("Date:", align=WD_ALIGN_PARAGRAPH.LEFT, size=11, space_after=0)

    insert_page_break_before()

    # =========================================================================
    # 4. ACKNOWLEDGEMENT
    # =========================================================================
    insert_p_before("ACKNOWLEDGEMENT", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_before=4, space_after=14)

    ack_p1 = (
        "It is with immense pleasure that we would like to express my indebted gratitude to my project guide "
        "Dr. Jalaiah Saikam, Department of Computer Science And Engineering, who has guided us a lot and "
        "encouraged in every step of the project work. His valuable moral support and guidance throughout the project helped me to a greater extent."
    )
    insert_p_before(ack_p1, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11.5, space_after=10, line_spacing=1.2)

    ack_p2 = (
        "We are grateful to Dr. A. Phani Sridhar HOD, for inspiring me all the way and for arranging all the "
        "facilities and resources needed for our project."
    )
    insert_p_before(ack_p2, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11.5, space_after=10, line_spacing=1.2)

    ack_p3 = (
        "We wish to thank our Dr. M.V. Rajesh, Associate Dean (School of Computing), Dr. A Ramesh, Pro Vice-Chancellor (Engineering & Sciences), "
        "Dr. S. Rama Sree, Pro Vice-Chancellor (Academics) and Dr. G. Suresh, Registrar, for their encouragement and support during the course of my project."
    )
    insert_p_before(ack_p3, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11.5, space_after=10, line_spacing=1.2)

    ack_p4 = (
        "We would like to extend my sincere thanks to Dr. M. B. Srinivas, Vice-Chancellor, Dr. M. Sreenivasa Reddy, Deputy Pro-Chancellor and Management, "
        "Aditya University for their unconditional support in providing me with the best infrastructural facilities and state-of-the-art laboratories during my project."
    )
    insert_p_before(ack_p4, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11.5, space_after=10, line_spacing=1.2)

    ack_p5 = (
        "Not to forget, Non-Teaching Staff and our Friends who have directly or indirectly supported me in completing this project on time"
    )
    insert_p_before(ack_p5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, size=11.5, space_after=20, line_spacing=1.2)

    insert_p_before("Student Name: Tanuku Ram Sai\nRegd. No: 24A95A0503", align=WD_ALIGN_PARAGRAPH.RIGHT, bold=True, size=11.5, space_after=0)

    insert_page_break_before()

    # =========================================================================
    # 5. LEARNING OBJECTIVES / INTERNSHIP OBJECTIVES
    # =========================================================================
    insert_p_before("Learning Objectives/Internship Objectives", align=WD_ALIGN_PARAGRAPH.LEFT, bold=True, size=14, space_before=6, space_after=12)
    insert_bullet_before("Internships are generally thought of to be reserved for college students looking to gain experience in a particular field. However, a wide array of people can benefit from Training Internships in order to receive real world experience and develop their skills.")
    insert_bullet_before("An objective for this position should emphasize the skills you already possess in the area and your interest in learning more")
    insert_bullet_before("Internships are utilized in a number of different career fields, including architecture, engineering, healthcare, economics, advertising and many more.")
    insert_bullet_before("Some internships are used to allow individuals to perform scientific research while others are specifically designed to allow people to gain first-hand experience working.")
    insert_bullet_before("Utilizing internships is a great way to build your resume and develop skills that can be emphasized in your resume for future jobs. When you are applying for a Training Internship, make sure to highlight any special skills or talents that can make you stand apart from the rest of the applicants so that you have an improved chance of landing the position.")

    insert_page_break_before()

    # =========================================================================
    # 6. INTERNSHIP COMPLETION CERTIFICATE
    # =========================================================================
    insert_p_before("INTERNSHIP COMPLETION CERTIFICATE", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=15, space_before=8, space_after=12)
    if os.path.exists(cert_img_path):
        insert_img_before(cert_img_path, width_in=5.8, space_before=8, space_after=8)
    else:
        insert_p_before("[ Official Internship Certificate Copy Attached — Tanuku Ram Sai (24A95A0503) ]\nIssued by: VedXlence Innovations Pvt. Ltd., Hyderabad\nAICTE Certified Upskilling & Project Implementation Program", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, italic=True, size=11)

    insert_page_break_before()

    # Save to updated paths
    out_file1 = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report_Updated.docx'
    out_file2 = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\24A95A0503_ShieldJob_AI_Internship_Report.docx'
    out_file_orig = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report.docx'

    doc.save(out_file1)
    print(f"Saved to: {out_file1} (size: {os.path.getsize(out_file1)} bytes)")
    doc.save(out_file2)
    print(f"Saved to: {out_file2} (size: {os.path.getsize(out_file2)} bytes)")

    try:
        doc.save(out_file_orig)
        print(f"Successfully overwritten: {out_file_orig}")
    except Exception as e:
        print(f"Note: {out_file_orig} could not be overwritten directly while open in Word.")

if __name__ == "__main__":
    perform_in_place_update()
