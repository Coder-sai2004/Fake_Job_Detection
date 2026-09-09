import docx
import os

path = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report_Updated.docx'
doc = docx.Document(path)
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"File size: {os.path.getsize(path)} bytes")

print("\n=== FIRST 25 PARAGRAPHS (PRELIMINARY PAGES) ===")
for i in range(min(25, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    if p.text.strip():
        print(f"P{i:03d} [{p.style.name}]: {p.text[:90]}")

print("\n=== PARAGRAPHS AROUND TOC ===")
for i, p in enumerate(doc.paragraphs):
    if "TABLE OF CONTENTS" in p.text.upper() or "3. INTRODUCTION" in p.text.upper():
        print(f"P{i:03d} [{p.style.name}]: {p.text[:90]}")

print("\n=== ALL TABLES ===")
for i, t in enumerate(doc.tables):
    print(f"Table {i}: {len(t.rows)} rows x {len(t.columns)} cols -> Header: {[c.text.strip() for c in t.rows[0].cells]}")

print("\n=== COUNTING IMAGES IN DOCUMENT ===")
img_count = 0
for rel in doc.part.rels.values():
    if "image" in rel.target_ref:
        img_count += 1
print(f"Total embedded images in docx: {img_count}")
