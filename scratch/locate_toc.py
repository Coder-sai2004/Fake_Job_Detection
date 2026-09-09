import docx
import sys

path = r'C:\Users\Ram Sai\OneDrive\Desktop\PROJECTS\Fake_Job_Detection\scratch\temp_report.docx'
doc = docx.Document(path)

print("=== FINDING TABLE OF CONTENTS IN DOCX ===")
toc_p_idx = -1
for i, p in enumerate(doc.paragraphs):
    if "TABLE OF CONTENTS" in p.text.upper():
        toc_p_idx = i
        print(f"Found 'TABLE OF CONTENTS' at paragraph P{i:03d}: '{p.text}'")
        break

print(f"TOC paragraph index: {toc_p_idx}")
if toc_p_idx != -1:
    print("\nParagraphs right before TOC:")
    for j in range(max(0, toc_p_idx - 10), toc_p_idx):
        print(f"P{j:03d} [{doc.paragraphs[j].style.name}]: '{doc.paragraphs[j].text[:80]}'")

    print("\nParagraphs right after TOC:")
    for j in range(toc_p_idx, min(len(doc.paragraphs), toc_p_idx + 15)):
        print(f"P{j:03d} [{doc.paragraphs[j].style.name}]: '{doc.paragraphs[j].text[:80]}'")
