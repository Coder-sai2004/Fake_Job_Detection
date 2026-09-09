import os
import shutil
import pypdfium2
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

WORKSPACE_DOC_DIR = r"C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship"
ORIGINAL_REPORT_PATH = os.path.join(WORKSPACE_DOC_DIR, "ShieldJob_AI_Internship_Report.docx")
FRIEND_REPORT_PATH = os.path.join(WORKSPACE_DOC_DIR, "INTERNSHIP REPORT.docx")
CERTIFICATE_PDF_PATH = os.path.join(WORKSPACE_DOC_DIR, "Internship_Certificate_2026.pdf")
FORMATTED_REPORT_PATH = os.path.join(WORKSPACE_DOC_DIR, "ShieldJob_AI_Internship_Report_Formatted.docx")
FINAL_REPLACE_PATH = ORIGINAL_REPORT_PATH
BACKUP_REPORT_PATH = os.path.join(WORKSPACE_DOC_DIR, "ShieldJob_AI_Internship_Report_Original_Backup.docx")

# 1. Create a backup of the original report
if not os.path.exists(BACKUP_REPORT_PATH) and os.path.exists(ORIGINAL_REPORT_PATH):
    shutil.copy2(ORIGINAL_REPORT_PATH, BACKUP_REPORT_PATH)
    print("Backup created at:", BACKUP_REPORT_PATH)

# 2. Extract Certificate PDF to high-res image
MEDIA_DIR = "report_assets"
os.makedirs(MEDIA_DIR, exist_ok=True)

pdf = pypdfium2.PdfDocument(CERTIFICATE_PDF_PATH)
cert_page = pdf[0]
# 300 DPI high-res rendering
cert_img = cert_page.render(scale=300/72).to_pil()
cert_img_path = os.path.join(MEDIA_DIR, "internship_certificate.png")
cert_img.save(cert_img_path, "PNG")
print("Rendered certificate to:", cert_img_path, "Size:", cert_img.size)

# 3. Extract friend doc logos
doc_friend = docx.Document(FRIEND_REPORT_PATH)
friend_logo_emblem = os.path.join(MEDIA_DIR, "aditya_emblem.jpg")
friend_logo_banner = os.path.join(MEDIA_DIR, "aditya_banner.png")

for rel in doc_friend.part.rels.values():
    if 'image1.jpeg' in rel.target_ref:
        with open(friend_logo_emblem, 'wb') as f:
            f.write(rel.target_part.blob)
    if 'image2.png' in rel.target_ref or 'image3.png' in rel.target_ref:
        with open(friend_logo_banner, 'wb') as f:
            f.write(rel.target_part.blob)

print("Extracted friend logos:", os.path.exists(friend_logo_emblem), os.path.exists(friend_logo_banner))

# 4. Extract user doc images
doc_user = docx.Document(ORIGINAL_REPORT_PATH)
user_images = {}
for rel in doc_user.part.rels.values():
    if 'image' in rel.target_ref:
        fname = os.path.basename(rel.target_ref)
        target_path = os.path.join(MEDIA_DIR, fname)
        with open(target_path, 'wb') as f:
            f.write(rel.target_part.blob)
        user_images[fname] = target_path

print(f"Extracted {len(user_images)} images from user report.")
