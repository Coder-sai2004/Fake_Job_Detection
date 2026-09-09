import docx
from docx.oxml import parse_xml
import xml.etree.ElementTree as ET

friend_path = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\INTERNSHIP REPORT.docx'
user_path = r'C:\Users\Ram Sai\OneDrive\Documents\24A95A0503_T.Ram Sai_internship\ShieldJob_AI_Internship_Report.docx'

doc_f = docx.Document(friend_path)
doc_u = docx.Document(user_path)

def inspect_images(doc, label):
    print(f"\n==================== {label} IMAGES & SHAPES ====================")
    for i, p in enumerate(doc.paragraphs):
        # find blip or drawings
        xml_str = p._p.xml
        if '<a:blip' in xml_str or '<w:drawing>' in xml_str or '<v:shape' in xml_str:
            print(f"Paragraph {i:03d} (align={p.alignment}): text='{p.text.strip()}'")
            # extract cx, cy if any
            root = ET.fromstring(xml_str)
            for elem in root.iter():
                if elem.tag.endswith('extent'):
                    cx = elem.attrib.get('cx')
                    cy = elem.attrib.get('cy')
                    if cx and cy:
                        cx_in = int(cx) / 914400
                        cy_in = int(cy) / 914400
                        print(f"   Image Extent: {cx_in:.2f} in x {cy_in:.2f} in (width x height)")
                if elem.tag.endswith('blip'):
                    embed = elem.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
                    if embed:
                        target = doc.part.rels[embed].target_ref
                        print(f"   Image Target: {target}")

inspect_images(doc_f, "FRIEND")
inspect_images(doc_u, "USER")
